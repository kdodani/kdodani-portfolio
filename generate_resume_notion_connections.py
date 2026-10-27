from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
# Background: #faf9f6  (warm off-white — body bg)
# Headings:   #1c1917  (stone-900 — near-black warm)
# Body text:  #44403c  (stone-700)
# Muted:      #78716c  (stone-500)
# Accent:     #6D28D9  (--accent-hero-1 purple)
# Mid-accent: #5B21B6  (--accent-strong)
# Sub-accent: #4F46E5  (--accent-hero-2 indigo)
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=8, space_after=2, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=5, after=1)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=1, after=1)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  US work authorization  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Connected Workflows  •  Partner Integrations  •  "
    "Admin Trust & Permissions  •  AI in the Product"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product manager who gets a connection adopted inside the tool people already use, then scales it. "
    "At WorkSafeBC, redirected clinical reporting from another web form to a direct API inside the EMR "
    "physicians would not leave — a reusable gateway, partner onboarding, and a roadmap covering about 70% "
    "of the vendor ecosystem — and cut a separate high-volume submission flow from 11 screens to 5 "
    "(paper share 56% to 19%). Makes tradeoffs across user experience, trust, technical complexity, and "
    "speed, and uses AI both to prototype and to ship.",
    size=9.5,
    space_before=4,
    space_after=4,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager", "WorkSafeBC", "Apr 2024 – Present")

sublabel(doc, "Connected systems")
bullet(doc,
    "Discovery showed physicians would not leave their EMR for a portal or another web form. Redirected "
    "the program to a direct API installed inside OSCAR, the open-source EMR used as the beachhead, then "
    "sequenced commercial vendors from that proof.")
bullet(doc,
    "Owned the integration decisions with architects: JSON for the March 2025 proof of concept, HL7 FHIR "
    "documented as the upgrade path, and Azure APIM as a reusable gateway instead of a point-to-point "
    "integration per vendor. Engineering implemented OAuth 2.0 / mTLS on that gateway.")
bullet(doc,
    "Secured $80K in Innovation Council funding. Led outreach to 50 EMR vendors and defined onboarding "
    "needs — API documentation and a Postman schema, with engineering — for a roadmap covering about 70% "
    "of the provider ecosystem.")

sublabel(doc, "One strategy across segments")
bullet(doc,
    "25 of 3,795 organizations generated half of the remaining paper (~280K line items). Set one channel "
    "strategy for those segments — redesigned portal, bulk vendor upload, and direct EMR — so differing "
    "needs became a roadmap rather than a list of one-off requests.")
bullet(doc,
    "Cut the submission journey from 11 screens to 5. Paper share of line items fell from 56% to 19%; "
    "digital channel share rose from 39% to 55%; support calls fell 30%; NPS rose 25; about $200K annual "
    "savings. Secured $580K to scale. Instrumented PostHog before launch and ran named-cohort dark, soft, "
    "and hard launches.")

sublabel(doc, "Trust and admin rollout")
bullet(doc,
    "Led the identity migration for 5,600 professionals across 500 organizations on a portal carrying "
    "62,000+ annual referrals — zero service disruptions, and 100% of active accounts transitioned on deadline.")
bullet(doc,
    "On an $80M program, 79% of 139K requests were always approved. With policy and risk partners, declined "
    "auto-approval so a human kept the decision, and shipped tracking, validation, and decision rationale.")
bullet(doc,
    "Prototyped Referral Radar, an LLM/RAG assistant that recommended the next referral and kept the "
    "clinician as the decision-maker. Placed 2nd of 36 at InnoFest 2025.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager (promoted from Business Analyst)", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Shipped NLP email classification into a live routing workflow, automating triage of incoming requests "
    "across multiple operational stakeholders.")
bullet(doc,
    "Led discovery for a mobile remote-inspections platform — field interviews, workflow mapping, and "
    "usability testing — and used that evidence to decide whether to invest in the build.")
bullet(doc,
    "Established the outcome metrics leadership used to prioritize retention: an 11.7% average annual churn "
    "rate, and an RFM segmentation model whose campaigns retained 78% of high-risk customers.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Designed a growth strategy for Open Primaries that cut bounce rate 37% while doubling subscriptions.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant — Growth & Analytics", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Built acquisition and forecasting systems for 75+ digital properties: organic performance ~150%, SEO "
    "forecast accuracy ~85%.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")
body(
    doc,
    "ProductCamp BC speaker (2025): “How Legacy Industries Can Adopt a Product Mindset.” Mentor, ProductBC.",
    size=9.5,
    space_before=2,
    space_after=2,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY

skills_row(doc, "Product",
    "Discovery to Roadmap  •  0→1 and Scale  •  Multi-segment Strategy  •  "
    "Tradeoffs across UX, Trust, Complexity, and Speed  •  Instrumentation and Launch")
skills_row(doc, "Platform and trust",
    "Partner Integrations  •  API Onboarding  •  Identity and Admin Rollout  •  "
    "Permissions and Human Decision Boundaries  •  External Systems of Record")
skills_row(doc, "Technical and AI",
    "REST APIs  •  Azure APIM  •  OAuth 2.0 as a Product Surface  •  SQL  •  Python  •  "
    "PostHog  •  LLM / RAG Product Design")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Notion_Connections.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
