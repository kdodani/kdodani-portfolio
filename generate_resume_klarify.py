from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins (match tech baseline) ────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)
    section.page_width    = Inches(8.5)
    section.page_height   = Inches(11)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def set_run_font(run, name="Calibri"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1.0

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=6, space_after=1, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    set_run_font(run)
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=3, after=0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    set_run_font(r1)
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    set_run_font(r2)
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=1, space_after=0):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    set_run_font(run)
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.18), hanging=Inches(0.15)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=0, after=0)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    set_run_font(run)
    return p

def body(doc, text, size=9.5, color=BODY, space_before=1, space_after=1, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    set_run_font(run)
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:cs"), "Calibri")
    rPr.append(rFonts)

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    set_run_font(run)
    return run

# Tighten the built-in list style so Word does not add its own gap.
list_style = doc.styles['List Bullet']
list_style.font.size = Pt(9.5)
list_style.font.name = "Calibri"
list_style.paragraph_format.space_before = Pt(0)
list_style.paragraph_format.space_after = Pt(0)
list_style.paragraph_format.line_spacing = 1.0

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=1)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK
set_run_font(r)

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=1)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  US work authorization  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=1, after=1)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)
set_run_font(rt)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=0, after=2)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Clinician Adoption  •  AI-Assisted Workflows  •  Regulated Healthcare  •  Ships with Cursor"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
set_run_font(rp)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product manager since 2021. Owns clinician-facing products in regulated healthcare when the brief is incomplete: "
    "notes, invoices, and treatment admin for physicians and physiotherapists who have no time to learn software. "
    "The simplest version is the one that ships.",
    size=9.5,
    space_before=2,
    space_after=1,
)
body(
    doc,
    "Speaks with those providers and closes the loop into the roadmap. Short cycles: Referral Radar in a 2-week sprint; "
    "the provider-invoicing rebuild and InfoX, an EMR proof of concept, each in under 3 months. Opens pull requests "
    "that reach production. Earlier, built organic-search acquisition systems for getting found.",
    size=9.5,
    space_before=0,
    space_after=2,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager — Adoption & Clinical Workflows", "WorkSafeBC", "Apr 2024 – Present")

sublabel(doc, "Adoption, experiments, and support")
bullet(doc,
    "Spoke with physicians, physiotherapists, and office admins, and closed a feedback loop into the product. "
    "The physio tool was used by office admins, not the therapists answering the survey, so later research asked "
    "who actually did the paperwork. Post-launch, users were missing a control that already existed — visibility, "
    "not a missing feature. Fixed the workflow, not the ticket queue; support calls −30%; NPS +25.")
bullet(doc,
    "Instrumented PostHog before launch and ran dark → soft → full releases so 322K existing digital lines a year "
    "stayed live while a new segment activated. Migrated 5,600 healthcare professionals across 500 organizations "
    "(62K+ annual transactions) with 0 disruptions — the constraint was behavior, not the identity stack.")

sublabel(doc, "Simplest workflow that gets used")
bullet(doc,
    "In under 3 months, rebuilt high-volume provider invoicing — the billing admin clinics finish after the visit — "
    "after 25 of 3,795 organizations drove 50% of paper volume (~280K line items). The constraint was the product "
    "(one claim per submit, 11 screens, line caps), not training. Cut the path to 5 screens. Product share "
    "39% → 55%; paper share 56% → 19%; digital adoption to 60% within six months; ~$200K annual savings; $580K to scale it.")

sublabel(doc, "Roadmap changed by customers")
bullet(doc,
    "Provider conversations killed a side portal: physicians will not leave the EMR. Led InfoX, a proof of concept "
    "in under 3 months, so the clinical note could be sent from OSCAR instead. POC validated March 2025; $80K funding; "
    "vendor roadmap covering ~70% of the ecosystem. JSON so validation was weeks, not a FHIR program; FHIR stayed "
    "the upgrade path; OAuth 2.0 / mTLS on the clinical data channel.")

sublabel(doc, "Clinical notes, judgment, and privacy")
bullet(doc,
    "Treated the clinical note as the product. Form 8/11 notes — written by physicians after the visit — were truncated "
    "by a 2001 billing spec (52,000 faxes a year). A shorter workflow and front-end validation improved the note that "
    "reached review, instead of asking clinicians to type around the limit.")
bullet(doc,
    "Refused auto-approve on physiotherapy treatment extensions after 139K requests showed a 79% approval rate. "
    "Speed would have rubber-stamped an $80M workflow of admin that follows the session. Shipped tracking and "
    "error validation, and kept clinical review in the product.")
bullet(doc,
    "In a 2-week sprint, built Referral Radar, an LLM/RAG prototype (2nd of 36, InnoFest 2025), over 20–75 documents "
    "per case — intake notes, medical reports, and counselling plans. Evidence traces and a visual map of issues, "
    "interventions, and outcomes; humans kept the decision. Production records were blocked, so the sprint used "
    "synthetic data checked with subject-matter experts.")

sublabel(doc, "Ships code")
bullet(doc,
    "Opened engineer-reviewed pull requests on the WorkSafeBC provider platform that merged to production. Built and "
    "shipped this portfolio (Next.js, React, Tailwind on Vercel) in Cursor.")

doc.add_page_break()

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager (promoted from Business Analyst)", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Shipped NLP email classification into a live routing workflow, cutting manual inbox triage — production AI, not a demo.")
bullet(doc,
    "Led discovery for a mobile inspection platform through customer interviews, workflow mapping, prototyping, and "
    "usability testing before engineering investment. Turned field pain into requirements and a roadmap with engineering, "
    "field users, and design.")
bullet(doc,
    "Built membership, churn, and engagement data products that identified 11.7% annual churn. RFM segmentation "
    "retained 78% of high-risk customers.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Designed a growth acquisition system for Open Primaries that cut bounce rate 37% while doubling subscriptions. "
    "Built product and market strategies for nonprofit and clean energy clients.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant — Growth & Analytics", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Built organic acquisition systems for 75+ digital properties, increasing organic performance ~150%, and a "
    "forecasting model that predicted SEO performance with ~85% accuracy.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT, space_before=0, space_after=0)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=1)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    set_run_font(rl)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY
    set_run_font(rs)

skills_row(doc, "Product",
    "Discovery  •  Roadmaps  •  Adoption  •  Provider interviews  •  Feedback loops  •  "
    "0→1  •  Short cycles  •  Support-to-product")
skills_row(doc, "Clinical",
    "Clinical notes  •  Physiotherapy admin  •  Provider invoicing  •  Human-in-the-loop decisions  •  "
    "Privacy-constrained data  •  Regulated healthcare")
skills_row(doc, "AI & building",
    "LLM / RAG  •  NLP (shipped)  •  Cursor  •  GitHub Copilot  •  Lovable  •  n8n  •  "
    "Evidence traces  •  Next.js  •  React")
skills_row(doc, "Growth",
    "Organic acquisition  •  Activation  •  PostHog  •  Funnel analysis  •  Segmentation  •  SQL  •  Python")

add_horizontal_rule(doc)
heading(doc, "Leadership", space_before=6)

sublabel(doc, "Speaker — ProductCamp BC (2025)", italic=False, size=9.5, color=DARK, space_before=1)
body(doc,
    '"How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery."',
    size=9.5, italic=True, space_before=0, space_after=1)

sublabel(doc, "Mentor — ProductBC", italic=False, size=9.5, color=DARK, space_before=2)
body(doc,
    "Mentor aspiring product managers in discovery, strategy, and career development.",
    size=9.5, space_before=0, space_after=0)

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Klarify.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
