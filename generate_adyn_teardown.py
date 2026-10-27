from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

DARK = RGBColor(0x1c, 0x19, 0x17)
ACCENT = RGBColor(0x6D, 0x28, 0xD9)
BODY = RGBColor(0x44, 0x40, 0x3c)
LIGHT = RGBColor(0x78, 0x71, 0x6c)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def spacing(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)


def rule():
    p = doc.add_paragraph()
    spacing(p, 1, 1)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "6D28D9")
    pBdr.append(b)
    pPr.append(pBdr)


def head(text, before=5):
    p = doc.add_paragraph()
    spacing(p, before, 2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = ACCENT


def para(parts, size=9, before=0, after=2):
    """parts: list of (text, bold) tuples or a string."""
    p = doc.add_paragraph()
    spacing(p, before, after)
    if isinstance(parts, str):
        parts = [(parts, False)]
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = DARK if bold else BODY
    return p


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_text(cell, parts, size=8.5):
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    spacing(p, 1, 1)
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = DARK if bold else BODY


# ── Title ────────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
spacing(p, 0, 0)
r = p.add_run("adyn Product Experience Teardown")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = DARK

p = doc.add_paragraph()
spacing(p, 0, 2)
r = p.add_run(
    "A first-week look at The Birth Control Test journey, from Khushboo Dodani. Built only from public information "
    "(adyn.com, FAQ, customer write-ups, Nurx / Hers / Natural Cycles). These are hypotheses to test, not claims."
)
r.italic = True
r.font.size = Pt(8.5)
r.font.color.rgb = LIGHT
rule()

# ── Journey ──────────────────────────────────────────────────────────────────
head("The journey as published", before=2)
para(
    "Quiz and purchase ($369, no insurance, HSA/FSA accepted)  >  kit arrives in 2-7 days  >  hormone sample must be "
    "taken on cycle day 3, saliva DNA any time  >  lab results in 2-6 weeks (avg. 3)  >  online report  >  book a "
    "25-minute provider visit  >  prescription (mail-order option)  >  retest hormones over time.",
    size=9,
)
para(
    [("Comparators: ", True),
     ("Nurx and Hers get a user to a prescription decision via a questionnaire and provider review (Hers replies "
      "reportedly within 24 hours; Nurx $28 consult + about $15 per month for medication). Natural Cycles sells daily "
      "engagement through a subscription app. adyn is the only one that sells precision, and it pays for that "
      "precision with time.", False)],
    size=9,
)

# ── Observations ─────────────────────────────────────────────────────────────
head("Three observations, hypotheses, and what I would measure")

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
for c, t in zip(hdr, ["What I see", "Hypothesis to test", "Metric"]):
    cell_text(c, [(t, True)], size=8.5)
    shade(c, "EDE9FE")

rows = [
    (
        [("1. The wait is the product risk. ", True),
         ("Value arrives 3-5+ weeks after purchase, and the sample timing depends on the user's cycle. "
          "Questionnaire-based competitors reach a provider decision in roughly a day (per reviews).", False)],
        [("A \"while you wait\" experience (kit status, day-3 cycle reminders, plain-language prep for the visit) "
          "lowers drop-off between purchase and sample return.", False)],
        [("Purchase to valid-sample-returned rate; first-pass sample validity; days to valid sample.", False)],
    ),
    (
        [("2. The report is rich but dense. ", True),
         ("Customers describe \"a TON of information\" with footnotes, and skeptics question the evidence. "
          "Trust is a core conversion lever for a genetics-based claim.", False)],
        [("Leading with \"your top matches and why, in your own words\" and tagging evidence strength raises "
          "confidence and speeds the decision.", False)],
        [("Report open to visit booked; time from results to booking; visit-to-prescription rate.", False)],
    ),
    (
        [("3. Results-to-visit is a handoff. ", True),
         ("Booking sits after the report; one public customer account describes being unable to schedule while "
          "the site was down. One-time $369 purchase, with retesting recommended.", False)],
        [("Booking the visit inline from the report (and offering a doctor-share summary) increases conversion, "
          "and a retest prompt tied to life events builds the longitudinal relationship adyn is betting on.", False)],
        [("Results-to-fulfillment rate; retest rate at 12 months; side-effect-driven method-switch rate "
          "at 3 and 6 months.", False)],
    ),
]
for a, b, c in rows:
    cells = table.add_row().cells
    cell_text(cells[0], a)
    cell_text(cells[1], b)
    cell_text(cells[2], c)

for row in table.rows:
    row.cells[0].width = Cm(7.0)
    row.cells[1].width = Cm(6.4)
    row.cells[2].width = Cm(4.4)

# ── Concept ──────────────────────────────────────────────────────────────────
head("Concept to prototype first: \"Results Home\"")
para(
    [("Before results: ", True),
     ("a four-step tracker (kit, sample, lab, results) with cycle-aware reminders and a one-tap video call for "
      "collection help. ", False),
     ("At results: ", True),
     ("top three matches with one-line reasons tied to the goals the user entered, an evidence-strength tag on each "
      "claim, the full report one tap deeper, and the visit booking calendar on the same screen. ", False),
     ("After the visit: ", True),
     ("a shareable summary for the user's own doctor and a retest reminder tied to life changes.", False)],
    size=9,
)
para(
    "Test plan: moderated usability sessions with 8-10 current or prospective birth control users, including users "
    "across race, age, and gender identity, to match adyn's inclusion mission; then an A/B test of the results "
    "layout against the current report.",
    size=9,
)

# ── KPIs ─────────────────────────────────────────────────────────────────────
head("Proposed KPI set")
kt = doc.add_table(rows=1, cols=3)
kt.style = "Table Grid"
kt.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, t in zip(kt.rows[0].cells, ["Clinical", "Engagement", "Business"]):
    cell_text(c, [(t, True)])
    shade(c, "EDE9FE")
kcells = kt.add_row().cells
cell_text(kcells[0], [("Side-effect-driven switches at 3 and 6 months (north star). Visit completion. "
                       "Adverse-event reports.", False)])
cell_text(kcells[1], [("Purchase to valid sample. Report to visit booked. Retest rate. Report depth and return visits.",
                       False)])
cell_text(kcells[2], [("Quiz to purchase. Results to prescription fulfillment. Cost per completed journey. "
                       "NPS and referral rate.", False)])

rule()
para(
    [("Why this fits the role: ", True),
     ("it mirrors the first-month brief (UX analysis of consumer digital health, prototypes, priorities, metrics) "
      "and uses the same loop I run today: research, instrument, test, and decide on evidence. "
      "Khushboo Dodani  |  khushdodani@gmail.com  |  kdodani-portfolio.vercel.app", False)],
    size=8.5,
)

out = r"c:\Users\khush\Desktop\kdodani-portfolio\adyn_Experience_Teardown_Khushboo_Dodani.docx"
doc.save(out)
print("Saved:", out)
