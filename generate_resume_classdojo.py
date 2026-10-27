from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=5, space_after=2, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=4, after=1)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=0, after=1)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY
    return p

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER — 0→1 PRODUCTS & GROWTH")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "0→1 Product  •  Customer Research  •  Adoption & Monetization  •  Fast, Lean Execution"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product manager with 5 years in product who takes ideas from zero to adoption: talk to users, find the "
    "real constraint in the data, prototype with AI tools before engineering spends a dollar, then grow usage. "
    "Built an LLM recommendation product 0→1 (2nd of 36 at InnoFest 2025) and validated an API-native "
    "integration POC that won $80K, using synthetic data and no-code tools when production access was blocked.",
    size=9.5,
    space_before=3,
    space_after=2,
)
body(
    doc,
    "Grew product share 39% → 55% and cut paper share 56% → 19% (~3×) for 5,600 users by rebuilding the "
    "flow for the <1% of customers driving half of volume; won $580K to scale it. Earlier: doubled "
    "subscriptions on a consumer product (Open Primaries) and grew organic traffic ~150% across 75+ consumer "
    "sites. BBA in Marketing; Master of Business Analytics.",
    size=9.5,
    space_before=0,
    space_after=3,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager — 0→1 Products, Growth & Research", "WorkSafeBC", "Apr 2024 – Present")
body(
    doc,
    "Own products used by 5,600 users across 500 organizations (62K+ annual transactions). Validate with "
    "customers, data, and prototypes before engineering invests.",
    size=9, color=LIGHT, space_before=0, space_after=3,
)

sublabel(doc, "0→1 — new products from a blank page")
bullet(doc,
    "Built Referral Radar 0→1, an LLM/RAG recommendation product that ranks next-best matches from 20–75 "
    "unstructured documents per case, with evidence traces and natural-language search. Rejected a ranked "
    "list with no “why” after user feedback; humans keep the decision. 2nd of 36 at InnoFest 2025.")
bullet(doc,
    "When production data was blocked, got SME-validated synthetic data and prototyped with Lovable, GitHub "
    "Copilot, and n8n, proving the concept with no engineering spend.")
bullet(doc,
    "Led 0→1 discovery for an API-native integration after research showed users would not leave their system "
    "of record for a side application. POC validated March 2025; $80K funding; roadmap covering ~70% of the "
    "ecosystem.")

sublabel(doc, "Customer research — qualitative signal into product decisions", space_before=2)
bullet(doc,
    "Ran interviews, workflow mapping, and prototype testing with the people doing the work. In a second round "
    "removed a control from a prototype to see whether users missed it, ignored it, or created demand.")
bullet(doc,
    "Segmented users by workflow role (admin submits, practitioner delivers); behavior predicted adoption better "
    "than organization type. Killed a “ping when stuck” feature after aging data showed most requests resolve "
    "within 3 days.")

sublabel(doc, "Growth — finding product-market fit inside a segment", space_before=2)
bullet(doc,
    "Analyzed 3.7M line items: 322K lines/year already converting, so PMF for one segment. Pareto: 25 of 3,795 "
    "organizations (<1%) generated 50% of remaining volume (~280K lines). Constraint was a capacity ceiling, not "
    "training. Focused product and GTM on that cohort, not a broad campaign.")
bullet(doc,
    "Rebuilt the high-volume flow (multi-claim, in-product validation, 11 screens → 5). Product share 39% → 55%; "
    "paper 56% → 19% (~3×); digital adoption to 60% in six months; support calls −30%; NPS +25; ~$200K annual "
    "operational savings. Secured $580K in executive funding to scale it.")
bullet(doc,
    "Wrote the PostHog spec before launch (conversion, drop-off, v1 vs v2). Staged dark → soft → full launch kept "
    "322K existing lines/year uninterrupted. Month one showed near-zero adoption: users missed the duplicate-line "
    "control. Fixed visibility and GTM, not scope. Split messaging by segment after one narrative failed to "
    "convert every cohort.")

sublabel(doc, "Distribution & network adoption", space_before=2)
bullet(doc,
    "One surface could not convert every segment, so built three distribution paths: rebuilt portal "
    "(high-volume), partner/vendor channel (bottom-up demand), and in-workflow API.")
bullet(doc,
    "Moved 5,600 users across 500 organizations through an identity cutover run as an activation campaign: "
    "28-day multi-channel treatments, with 500 org admins as multipliers instead of 5,600 individuals. "
    "100% of active accounts transitioned; 0 service disruptions.")
bullet(doc,
    "On an $80M workflow, analyzed 139K requests (79% always approved), killed auto-approve to protect user "
    "trust, and shipped tracking, decision rationale, and error validation instead.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(
    doc,
    "Product Manager — Retention & Revenue (promoted from BA)",
    "Technical Safety BC",
    "Oct 2021 – Apr 2024",
)
bullet(doc,
    "Built lifetime-value and churn views across ~102,000 permits (Energy Utility ~$7,400 per permit). "
    "RFM-targeted outreach vs. the existing call strategy, same two weeks and measures set first: $45,742 "
    "collected vs. $25,597 (78% more).")
bullet(doc,
    "Designed a three-step reactivation sequence (introduction, firmer follow-up, regulatory language) for "
    "135 expired clients (~CAD 46.7K): 23% paid from the first email.")
bullet(doc,
    "Shipped NLP classification into a live routing workflow, matching inbound requests to the right queue by "
    "intent and urgency.")
bullet(doc,
    "Led discovery for a mobile inspection platform (interviews, workflow mapping, prototyping, usability "
    "testing) and held engineering spend until the workflow was validated.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Designed and optimized consumer-facing landing pages and acquisition funnels for Open Primaries, cutting "
    "bounce rate 37% and doubling subscriptions. Built product and market strategies for nonprofit and "
    "clean-energy clients through customer research and segmentation.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant — Growth & Analytics", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Owned organic acquisition for 75+ consumer-facing digital properties, lifting organic performance ~150%; "
    "built a forecasting model predicting SEO performance with ~85% accuracy so channel investment followed "
    "projected return.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia — Sauder School of Business", "")
body(
    doc,
    "Customer analytics  •  Data-driven marketing  •  Discrete choice  •  Forecasting  •  Python  •  SQL",
    size=9, color=LIGHT, space_before=0, space_after=1,
)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")
body(
    doc,
    "Consumer behavior, segmentation, and go-to-market.",
    size=9, color=LIGHT, space_before=0, space_after=1,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")
body(
    doc,
    "Speaker, ProductCamp BC (2025) — “How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery.”   "
    "Mentor, ProductBC — product discovery, strategy, and career development.",
    size=9,
    space_before=1,
    space_after=2,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

skills_row(doc, "Product",
    "0→1 Discovery  •  Product Strategy  •  Prototyping  •  Roadmaps  •  Staged Launches  •  "
    "Activation & Onboarding  •  Distribution")
skills_row(doc, "Research & analytics",
    "User Interviews  •  Usability Testing  •  Journey Mapping  •  Funnel & Cohort Analysis  •  A/B Testing  •  "
    "PostHog  •  SQL  •  Python  •  Forecasting")
skills_row(doc, "Growth & revenue",
    "Segmentation  •  Pareto  •  RFM  •  Lifetime Value  •  Retention  •  Reactivation  •  "
    "Acquisition Channels  •  SEO")
skills_row(doc, "AI",
    "LLM / RAG Prototyping  •  NLP (shipped)  •  Human-in-the-Loop Design  •  Lovable  •  GitHub Copilot  •  n8n")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_ClassDojo.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
