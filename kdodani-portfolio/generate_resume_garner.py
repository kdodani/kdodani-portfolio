from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
# Background: #faf9f6  (warm off-white — body bg)
# Headings:   #1c1917  (stone-900 — near-black warm)
# Body text:  #44403c  (stone-700)
# Muted:      #78716c  (stone-500)
# Accent:     #6D28D9  (--accent-hero-1 purple)
# Mid-accent: #5B21B6  (--accent-strong)
# Sub-accent: #4F46E5  (--accent-hero-2 indigo)
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=8, space_after=2, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=5, after=1)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=1, after=1)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "Relocating to New York City  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "B2B Client Platforms  •  External Portal + Internal Systems  •  Revenue & Eligibility  •  Data Products  •  AI-Enabled Operations"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product Manager with 5 years owning B2B healthcare platforms — both the external client-facing portal and "
    "the internal systems used by operations, implementations, and finance-adjacent teams. Experience covering "
    "invoicing, eligibility rules, and billing workflows, plus data products on membership and revenue. Strong "
    "systems thinker; designs for edge cases and failure modes. Ships 0-to-1 and scaled products with engineering "
    "and operations. Uses AI daily to analyze, prototype, and accelerate discovery. Master of Business Analytics; "
    "BBA in Marketing.",
    size=9.5,
    space_before=4,
    space_after=4,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager — External Portal & Internal Systems", "WorkSafeBC", "Apr 2024 – Present")

sublabel(doc, "Client Portal Strategy & Roadmaps")
bullet(doc,
    "Own both the external client portal and the internal tools that make it run — billing, eligibility, claims, "
    "identity, and implementations — so client organizations can interact with benefits infrastructure while "
    "operations keep data and cost control intact.")
bullet(doc,
    "Own product strategy and roadmap for the external portal healthcare organizations use to interact with "
    "WorkSafeBC — invoices, referrals, and treatment requests — translating outcomes into sequenced delivery "
    "with engineering, operations, and program leadership.")
bullet(doc,
    "Digitized physio treatment-extension requests (replacing fax/phone) with tracking, error validation, and "
    "phased rollout across provider and case-manager groups — partnering with CMS product owners on an "
    "$80M physiotherapy spend workflow so care and billing keep moving.")
bullet(doc,
    "Designed the internal case-manager side of the same workflow so authorization decisions stay in human "
    "review rather than auto-approve — 79% of 139K historical extensions were always approved, which would "
    "have invited rubber-stamping an $80M spend if the product optimized for speed alone.")

sublabel(doc, "Invoicing, Billing Rules & Financial Workflows")
bullet(doc,
    "Built the provider invoicing product organizations use to bill WorkSafeBC; encoded billing and claims "
    "rules so invalid services, claims, and line items fail in-product instead of in back-office cleanup.")
bullet(doc,
    "Rebuilt high-volume invoicing after 25 of 3,795 organizations generated 50% of remaining paper volume "
    "(~280K blocked line items): cut submission from 11 screens to 5, shifted paper from 56% to 19%, and grew "
    "portal share from 39% to 55% — securing $580K in executive funding and ~$50K in ELAN savings.")
bullet(doc,
    "Exposed contract-eligibility through the provider invoice API so organizations only see services they "
    "are signed up for — an onboarding and eligibility control that reduced invalid submissions before they "
    "reached operations.")
bullet(doc,
    "Instrumented PostHog funnels before launch and ran named-cohort dark / soft / hard launches with "
    "Healthcare Operations so existing 322K digital lines/year were not disrupted while a new high-volume "
    "segment activated — adoption risk was visible within days, not after a big-bang cutover.")

sublabel(doc, "Internal Systems — Implementations, Integrations & Ops")
bullet(doc,
    "Partnered across product, engineering, operations, and finance-adjacent teams to define requirements, "
    "tradeoffs, and rollout strategies for mission-critical invoicing, identity, and eligibility workflows.")
bullet(doc,
    "Led a CIAM migration of 5,600 healthcare professionals across 500 organizations on a portal carrying "
    "62,000+ annual referrals: multi-channel outreach, administrator-led rollout, and helpdesk enablement — "
    "0 service disruptions; 100% of active accounts transitioned on deadline.")
bullet(doc,
    "Led 0-to-1 product discovery for EMR-at-point-of-care submission so physicians could send clinical "
    "documentation from OSCAR rather than a separate portal; secured $80K Innovation Council funding and a "
    "roadmap covering ~70% of the provider ecosystem; coordinated 50 EMR vendors plus TELUS, ForgeRock, and "
    "government partners from requirements through go-live.")

sublabel(doc, "Zero-to-One & AI-Enabled Products")
bullet(doc,
    "Built Referral Radar, an AI-assisted next-best-referral product using LLM agents and explainable "
    "recommendations so case managers could surface program eligibility buried across 20–75 claim documents — "
    "2nd place at WorkSafeBC InnoFest 2025.")
bullet(doc,
    "Unblocked the sprint when production claim data was unavailable by building a synthetic-data path "
    "validated with case managers — keeping humans as the decision-makers and AI as the retrieval layer.")
bullet(doc,
    "Used AI daily (Lovable, GitHub Copilot) to prototype workflows, validate concepts with users, and "
    "compress discovery-to-engineering timelines before development investment.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Built internal data products (membership, churn, and engagement dashboards plus an RFM model) tracking "
    "11.7% average annual churn — used by operations and leadership to prioritize retention (78% of high-risk "
    "customers retained) and protect recurring revenue.")
bullet(doc,
    "Led discovery for a mobile remote inspections platform used by field operations — contextual research, "
    "workflow mapping, usability testing, and prototyping before development investment.")
bullet(doc,
    "Partnered with engineering, field inspectors, designers, and operations leaders to translate operational "
    "pain points into product requirements, roadmap priorities, and implementation plans.")
bullet(doc,
    "Built NLP-powered email classification to automate routing and triage of incoming operational requests, "
    "reducing manual inbox processing burden.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Developed product and market strategies for nonprofit and clean energy organizations through customer "
    "research, market analysis, and strategic planning; designed a data-driven acquisition strategy for Open "
    "Primaries that reduced website bounce rate by 37% while doubling subscriptions.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Built analytics-driven growth strategies for 75+ client websites, increasing organic performance by "
    "approximately 150%, and a forecasting tool predicting SEO performance with approximately 85% accuracy.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")

sublabel(doc, "Speaker — ProductCamp BC (2025)", italic=False, size=9.5, color=DARK)
body(doc,
    '"How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery."',
    size=9.5, italic=True, space_before=0)

sublabel(doc, "Mentor — ProductBC", italic=False, size=9.5, color=DARK, space_before=4)
body(doc,
    "Mentor aspiring Product Managers in product discovery, strategy, and career development.",
    size=9.5, space_before=0)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY

skills_row(doc, "Product",
    "Product Strategy  •  Product Roadmaps  •  B2B Client Portals  •  0-to-1 Products  •  "
    "Change Management  •  Digital Adoption  •  Stakeholder Alignment  •  Executive Influence")
skills_row(doc, "Financial & Data Systems",
    "Invoicing  •  Billing Rules  •  Claims Rules  •  Contract Eligibility  •  Data Products  •  "
    "Revenue Analytics  •  Membership / Churn Dashboards  •  RFM Segmentation  •  PostHog")
skills_row(doc, "Healthcare B2B",
    "Client Portals  •  External + Internal Systems  •  Internal Ops Tools  •  Implementations  •  "
    "Integrations  •  Healthcare Operations  •  Regulated Healthcare")
skills_row(doc, "Technical",
    "REST APIs  •  EMR/EHR Integrations  •  CIAM/IAM  •  Multi-Vendor Delivery  •  SQL  •  Python  •  "
    "Azure  •  GitHub Copilot  •  Lovable")
skills_row(doc, "Research",
    "Journey Mapping  •  User Interviews  •  Workflow Analysis  •  Usability Testing  •  A/B Testing  •  "
    "Prototyping  •  Figma")
skills_row(doc, "AI",
    "Generative AI  •  LLM Applications  •  NLP  •  Intelligent Task Routing  •  AI-assisted Prototyping")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Garner_Health.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
