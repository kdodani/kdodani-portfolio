from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Document() may start with an empty paragraph. Drop it so the header sits at the margin.
if doc.paragraphs:
    _blank = doc.paragraphs[0]._element
    _blank.getparent().remove(_blank)

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.27)
    section.bottom_margin = Cm(1.27)
    section.left_margin   = Cm(1.6)
    section.right_margin  = Cm(1.6)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.5)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0

bullet_style = doc.styles["List Bullet"]
bullet_style.font.name = "Calibri"
bullet_style.font.size = Pt(9.5)
bullet_style.paragraph_format.space_before = Pt(0)
bullet_style.paragraph_format.space_after = Pt(1)
bullet_style.paragraph_format.line_spacing = 1.0

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=6, space_after=1, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=4, after=0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.15), hanging=Inches(0.15)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=0, after=1)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER  |  HEALTH PLATFORMS, AI & DATA INTEGRATION")
rt.bold = True
rt.font.size = Pt(11)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Data Integration  •  Workflow Launches  •  LLM / RAG Prototyping  •  Product Analytics"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product Manager with 5 years of experience building and improving technical products across healthcare "
    "and operational platforms, including product discovery, production launches, and AI prototyping. Led "
    "healthcare data integration discovery, shipped a high-volume invoicing workflow from 11 screens to 5, "
    "and developed an LLM/RAG prototype that surfaces evidence-backed recommendations with human oversight. "
    "Combines user research, product analytics, API integration strategy, and cross-functional execution to "
    "improve data quality, adoption, and reliability in regulated environments. Holds a Master of Business Analytics.",
    size=9.5,
    space_before=4,
    space_after=4,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager — Health Data & Workflows", "WorkSafeBC", "Apr 2024 – Present")

sublabel(doc, "Health data integration")
bullet(doc,
    "Led discovery so clinicians could submit clinical documentation from OSCAR EMR instead of a side portal, "
    "after interviews showed they would not leave their system of record. This was a validated POC (March 2025), "
    "not a production launch. Innovation Council approved $80K for this POC — a separate approval from the "
    "$580K invoicing rebuild. Roadmap covered ~70% of the provider ecosystem, starting with OSCAR before paid "
    "EMR vendors.")
bullet(doc,
    "Set the integration tradeoff with engineering and solution architects: a JSON payload so the POC could be "
    "validated in weeks, with HL7 FHIR documented as the later upgrade path. Engineering implemented the "
    "reusable gateway (Azure API Management rather than a connection per vendor) and OAuth 2.0 / mTLS on the "
    "data channel.")

sublabel(doc, "Invoice workflow — shipped")
bullet(doc,
    "Found that 25 of 3,795 organizations drove 50% of remaining paper (~280K line items). The constraint was "
    "the product — one claim per submit, 11 screens, line caps — not training. Rebuilt high-volume submission "
    "with multi-claim entry and validation at submit.")
bullet(doc,
    "Redesigned high-volume invoice submission from 11 screens to 5. Mail and fax fell from 56% to 19% of "
    "invoice line items between 2023 and 2026. Separately, online portal share rose from 39% to 55% as "
    "high-volume providers moved onto the new version. Executive sponsors approved $580K for this rebuild "
    "before development. Early paper-volume reduction saved about $50K in processing cost to date.")
bullet(doc,
    "Instrumented PostHog before launch (conversion, drop-off, v1 vs v2, high-volume segment) and ran dark, "
    "soft, then full cohort launches so 322K existing digital lines per year stayed up while the new segment "
    "came on. The first month barely moved: providers missed the duplicate-line control. That was discoverability, "
    "not a missing feature, so the next iteration was onboarding rather than another build.")
bullet(doc,
    "On treatment requests, clinic administrators were the people submitting, while physiotherapists were the "
    "ones answering satisfaction surveys, so the scores mixed contract sentiment with the tool. Pointed discovery "
    "and form design at the person who submits.")

sublabel(doc, "Review, privacy, and access decisions")
bullet(doc,
    "On an $80M physiotherapy program (139K historical requests, 79% always approved, 27% resolved the same "
    "day), senior PMs and architects pushed to auto-approve. Rejected it: the approval rate was not evidence "
    "that review should stop. Shipped request tracking and error validation, and scoped decision rationale so "
    "a provider could see why a request was held. Partnered with the case-management product owner so the "
    "reviewer side and the submission side stayed aligned.")
bullet(doc,
    "Led the identity migration for 5,600 healthcare providers across 500 organizations on a portal carrying "
    "62K+ annual referrals. 0 service disruptions. 100% of active accounts had transitioned by the deadline. "
    "Emails were not on file, so the rollout went through organization admins over 28 days, using email, "
    "in-product banners, administrator-led and self-serve paths, and helpdesk enablement.")

sublabel(doc, "Generative AI — validated prototype")
bullet(doc,
    "Built Referral Radar, an LLM/RAG prototype for the case-review workflow. It surfaces evidence across "
    "20–75 unstructured documents per case and recommends a next program pathway, including mental-health "
    "pathways, with a source trace a reviewer can check. Humans kept the decision. Production claim records "
    "were blocked, so case managers validated synthetic records instead of waiting on protected health data. "
    "2nd of 36 at InnoFest 2025. Prototyped with Lovable, GitHub Copilot, and n8n before engineering investment. "
    "Validated prototype, not a production launch.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager (promoted from Business Analyst)", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Shipped NLP email classification into a live routing workflow so inbound requests left a shared inbox "
    "and landed in the right queue, cutting manual triage.")
bullet(doc,
    "Built membership, churn, and engagement data products that identified 11.7% average annual churn. "
    "RFM segmentation plus targeted engagement against baseline outreach retained 78% of high-risk customers.")
bullet(doc,
    "Led discovery for a mobile inspection platform through customer interviews, workflow mapping, prototyping, "
    "and usability testing before engineering investment. Partnered with engineering, field users, and design "
    "on requirements and the roadmap.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
sublabel(doc, "Consumer journeys")
bullet(doc,
    "Designed the acquisition journey for Open Primaries, a public website: bounce rate fell 37% and "
    "subscriptions doubled. Also built product and market strategies for nonprofit and clean-energy clients "
    "from customer research.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant — Growth & Analytics", "Convertus", "May 2018 – Jul 2019")
sublabel(doc, "Consumer web analytics")
bullet(doc,
    "Used behavioral and search analytics to grow organic performance about 150% across 75+ consumer-facing "
    "websites, and built a forecasting model that predicted SEO performance with about 85% accuracy.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")
body(doc, "Consumer behavior  •  Segmentation  •  Communication strategy", size=9, color=LIGHT)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")

sublabel(doc, "Speaker — ProductCamp BC (2025)", italic=False, size=9.5, color=DARK)
body(doc,
    '"How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery."',
    size=9.5, italic=True, space_before=0)

sublabel(doc, "Mentor — ProductBC", italic=False, size=9.5, color=DARK, space_before=4)
body(doc,
    "Mentor aspiring Product Managers in product discovery, strategy, and career development.",
    size=9.5, space_before=0)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

skills_row(doc, "Product",
    "Conception to Launch  •  0→1 / 1→n  •  Product Strategy  •  Roadmaps  •  Cross-functional Leadership  •  "
    "Prioritization  •  Staged Launches  •  Executive Funding")
skills_row(doc, "Health data",
    "Health Data Integration  •  EMR / EHR  •  HL7 FHIR (upgrade path)  •  REST APIs  •  Azure API Management  •  "
    "OAuth 2.0 / mTLS  •  CIAM  •  Validation at Submit")
skills_row(doc, "AI",
    "Generative AI in Workflows  •  LLM / RAG Interfaces  •  Evidence Traces  •  Human-in-the-Loop  •  "
    "NLP in a Live Workflow  •  GitHub Copilot  •  Lovable  •  n8n")
skills_row(doc, "Decisions",
    "Human Review Kept on Clinical Workflows  •  Privacy Constraints  •  Synthetic Data When Production Records Are Blocked  •  "
    "Identity & Access")
skills_row(doc, "Analytics",
    "PostHog  •  Funnel Instrumentation  •  Cohort Analysis  •  Experimentation  •  SQL  •  Python  •  "
    "Segmentation  •  RFM")
skills_row(doc, "Research",
    "User Interviews  •  Usability Testing  •  Workflow Analysis  •  Data-Collection Design  •  "
    "Survey Design  •  Prototyping")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Google_Health_Guardian.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
