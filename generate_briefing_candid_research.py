"""Candid Health interview research briefing (Word)."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor, Cm
from playbook_docx_lib import (
    new_doc, title_block, h1, h2, h3, body, bullet, labeled, say, note, DARK, MID, BODY, LIGHT,
)

doc = new_doc()
title_block(
    doc,
    "Khushboo Dodani  ·  interview briefing",
    "Candid Health — Research + Tell Me About Yourself",
    "Say the script first. Then use the Q&A. Public facts are sourced. WorkSafe items are your observations. No invented US RCM, CPT, or 837 experience.",
)

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "How to use this")
bullet(doc, "Part 1 is what you say on the recruiter / hiring-manager open. Practice it out loud once.")
bullet(doc, "Part 2 is company research — use it when they probe product, TCR, or unit economics.")
bullet(doc, "Part 3 is the WorkSafe → Candid map and the gap script.")
bullet(doc, "Vendor-reported metrics (95–99% TCR, customer collection rates) are labeled. Do not treat them as audited.")
note(doc, "The Candid product video you attached could not be transcribed in this environment (no local media tools). Parallels below use Candid’s published product thesis and the Oak HC/FT founder interview (same story the marketing video tells) plus the WorkSafe gaps you listed.")

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "Part 1 — Tell me about yourself")

h2(doc, "The arc (do not skip this)")
body(doc, "Gaps you saw in a legacy billing system → you chose to modernize that class of industry (ProductCamp) → two projects that prove detect + agent judgment → Candid is the closed loop you could not ship: auto-correct, learn from humans, enterprise visibility, touchless rate.")

h2(doc, "90-second script")
say(doc, [
    ("Bottom line", "I build products that take manual billing and claims work out of people’s queues — so providers get paid, and spend the time on care."),
    ("Gaps", "At WorkSafeBC I owned provider invoicing and the ops tools behind it. The platform could find errors but not auto-correct them. Every exception required a human, every time — the system did not learn. There was no stale-claim follow-up. Fee schedules and fee codes lived on a website; providers had to look them up to bill correctly. Large orgs like CBI Health only got siloed, location-level data — log in per site, no enterprise visibility, nothing that helped them decide and act."),
    ("Trajectory", "That is why I care about modernizing legacy healthcare and insurance. I spoke at ProductCamp BC in 2025 on how those industries adopt a product mindset. Candid is the next step on that same path."),
    ("PIS", "Provider Invoice Submission: time and dollars wasted, payments and treatment delayed. I encoded fee schedules, contracted rates, and service codes so invalid line items failed at submit — invalid submissions down 40%. We analyzed 3.7M line items, rebuilt throughput for the 25 orgs driving half the remaining paper, unlocked about 280K blocked lines. That is Candid’s detect layer. The gap I could not close: auto-correct and a feedback loop from ops touches into new rules."),
    ("Referral Radar", "An agent over 20–75 claim documents so a human can find program eligibility and still keep the decision. Same pattern as Candid’s agents — retrieval and judgment, not unsupervised spend. I have already refused to auto-approve an $80M workflow when 79% of 139K records were always approved."),
    ("Why Candid", "Candid’s product is the closed loop: encode payer rules, correct the claim before it goes out, learn when a human still has to touch it, and give the org one place to see and act. I want to do that work, in the San Francisco office, four days a week."),
], "75–90s")

h2(doc, "If they let you go two minutes — add these, do not restart")
bullet(doc, "Dashboards: TSBC membership/churn/RFM; PostHog on 322K digital lines/year before we cut over a high-volume segment. The JD asks for rigorous reporting. I already instrument launches.")
bullet(doc, "Change management: GTM and comms for PIS and the BCeID / CIAM migration (5,600 professionals, 0 downtime). The JD asks for rollouts to customers, ops, and leadership.")
bullet(doc, "Honest gap, one sentence: I do not have US commercial RCM — no CPT or 837 production work. I built detect-and-validate on a payer-adjacent billing system. Candid’s delta is auto-correct plus a learning loop. That is why I am here.")

h2(doc, "Phrases to hit / kill")
labeled(doc, "Hit", "Detect vs auto-correct. No learning loop. Fee codes on a website. Siloed enterprise data. First-pass / touchless. Human keeps judgment.")
labeled(doc, "Kill", "I’m passionate about healthcare. I’m not an RCM PM (as the opener). Cost per claim was $X (you do not have that number).")

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "Part 2 — Company research")

h2(doc, "How Candid makes RCM human-centered (specific examples)")
body(doc, "The JD says “complex payments ecosystems into a human-centered platform.” In their product that does not mean a prettier queue. It means humans stop doing keystrokes that a rule can do, and keep the work that is actually judgment.")
bullet(doc, "Multi-site service location. Legacy systems send every claim under one corporate address. Payers then treat the visit as out-of-network. Founders said customers had teams whose whole job was to open a queue and rewrite addresses. Candid overwrites the correct address from tax ID / site context. Human-centered = that team is gone; the biller is not a data-entry clerk. (Oak HC/FT interview, Nick Perry, Jul 2026.)")
bullet(doc, "Time-based therapy coding. Old billing systems often have no duration field. Humans read appointment notes to code units. Candid pulls appointment time from the integration layer and codes the claim. Human-centered = nobody re-reads a note to type a number the system already had.")
bullet(doc, "Correct before submit, not after denial. A bad claim that reaches the payer becomes a patient EOB that looks like a full denial, then a correction months later. Fixing upstream is human-centered for the patient, not only the biller. (Out-of-Pocket payment walkthrough; Sixth Street investment note.)")
bullet(doc, "Data that can be acted on. Legacy billers return CSVs. Candid’s pitch is reporting and a knowledge graph so finance and ops can see first-pass, fail reasons, and where humans still spend time — then turn that time into a new rule. Your CBI Health observation (siloed login, no enterprise action) is the opposite of this.")
bullet(doc, "Agents with a human residual. Eligibility, coding, appeals agents sit on the knowledge graph. The remaining ~5% is the human-centered surface: medical necessity, missing clinical facts, novel payer behavior — not address typing.")

h2(doc, "Primary users and pain points")
h3(doc, "Who buys / who uses")
bullet(doc, "Economic buyer: digital-health companies and, increasingly, multi-site outpatient groups, MSOs, and EMR partners (Talkiatry, Tia, Nourish, Cortica named publicly). They want billing to disappear or their existing team to run on better rails.")
bullet(doc, "Daily users: in-house billers and office managers (if hybrid); Candid’s own ops on the residual queue; customer finance/leadership on dashboards.")
bullet(doc, "Upstream data source, not the daily user: clinicians. Chart quality still decides whether a claim can be touchless.")
bullet(doc, "Engineering at digital-health customers: API-first ingest. Different from a fax-and-portal world — closer to your invoice API work.")

h3(doc, "Pains they are buying their way out of")
bullet(doc, "1,000+ payer rulebooks, formats, coding preferences, adjudication bugs. Rules change quarterly/annually (“adaptive adversaries”).")
bullet(doc, "Denials and rework. Cost-to-collect. Days in A/R. Collection rate leakage.")
bullet(doc, "Eligibility is still unsolved in the US (coverage yes; “does this plan cover these three therapy visits?” often no).")
bullet(doc, "Wrong payer routing (card logo ≠ claims address; Medicaid vs MCO).")
bullet(doc, "Black-box third-party billers: weak data back to the provider, hard to forecast or renegotiate contracts.")
bullet(doc, "Clean-claim vanity: a claim can be “clean” at the clearinghouse and still require a human and still pay wrong.")

h2(doc, "What “macro-structure of the business” means (JD)")
body(doc, "The JD is not asking you to recite the org chart. It is asking whether you think in how the company and the customer make or lose money on a claim — not in screens.")
labeled(doc, "Claim path", "Encounter → documentation → codes → claim (837) → clearinghouse → payer adjudication → ERA/835 → A/R / appeal → patient balance.")
labeled(doc, "Customer value (Nick Perry)", "Three levers only: collection rate (dollars you actually get), cost-to-collect (people + software), A/R time (days to cash). Every feature should move one of these.")
labeled(doc, "Candid’s P&L", "Revenue ≈ take-rate × customer collections (plus coding fees if they code). Cost-to-serve ≈ human minutes on non-touchless claims + impl. NRR grows when customer volume grows.")
labeled(doc, "Product surfaces", "Eligibility, coding, submit/edits, follow-up, denials/appeals, reporting. A PM who only owns “the portal homepage” is not at this layer.")
labeled(doc, "Segments", "API-first digital health (original beachhead) vs enterprise replace / white-label into EMR and legacy billing. Series D is a bet they can do both.")
labeled(doc, "Flywheel", "More claims → richer payer graph → higher TCR → lower cost-to-serve → better retention → more claims. Your WorkSafe system did not have this flywheel.")
note(doc, "Interview line: “I will not prioritize a UI polish that does not move collection rate, cost-to-collect, or A/R days.”")

h2(doc, "History and milestones (relevant to this role)")
bullet(doc, "2019–20: Founded by Nick Perry (Palantir healthcare) and Doug Proctor (Palantir; now COO), with Adam Reis. Y Combinator W20. Thesis: claims fail because they are submitted wrong, not because appeals software is weak.")
bullet(doc, "Early years: API-first RCM for venture-backed digital health that bills insurance. Word-of-mouth growth; little paid marketing (Fortune, 2026).")
bullet(doc, "2024: Series B ~$29M. Consecutive run-rate growth cited as 6x, then ~3–5x, then ~2.5x in later years (founder / Fortune).")
bullet(doc, "Feb 2025: Series C $52.5M led by Oak HC/FT; total ~$99.5M at that point.")
bullet(doc, "Jul 2026: Series D $120M led by Sixth Street Growth; Oak, 8VC, YC participate. Total raised past ~$219M. Valuation reported near $1B. Contracted run-rate +190% YoY; NRR described as very high (Fortune: close to 200%; Sixth Street: 180% for 2025 in some writeups — do not over-quote a single NRR number).")
bullet(doc, "Proof points used in fundraising: Talkiatry — ~40% less manual billing work, 98.3% of what payers owe; Nourish — 96.7% claims automatic, grew without adding billing staff. TCR commonly cited 95–99% / “as high as 97%.” Investor- and company-reported.")
bullet(doc, "Scale: on the order of 200+ provider organizations; billions in claims annually in investor materials. Team on the order of 150–200. Offices SF (HQ), NYC, Denver. 4 days in office — that is the culture, not a perk.")
bullet(doc, "Now: packaging automation/agents for customers who will not rip out 20-year-old billing systems; moving upmarket to traditional outpatient networks and EMR white-label. The PM JD exists because they need people who can own large surface area, dashboards, customer rollouts, and PM process while that happens.")
note(doc, "Why this matters in the interview: you are not joining a 10-person experiment. You are joining a company that is scaling product management (your JD: improve PM practices) while the core loop is already working at 95%+ TCR. The job is edges, enterprise, and the flywheel — not inventing touchless from zero.")

h2(doc, "What “unit economics of the business” means")
body(doc, "Unit economics = contribution profit on the repeating unit, before overhead. Here the unit is a claim (or a dollar of collections), not a user session.")
bullet(doc, "For the customer: dollars collected / dollars billed; cost to collect those dollars (biller FTEs + Candid fees); days in A/R. Perry’s three levers.")
bullet(doc, "For Candid: (take-rate × collections) − cost to process that volume (esp. human time on the non-touchless 1–5%). If TCR falls, Candid’s gross margin falls even if logo count is flat.")
bullet(doc, "Why the JD cares: a PM who ships a feature that adds ops minutes per claim is destroying unit economics. A PM who turns a repeated human touch into a rule is creating them.")
bullet(doc, "Your honest analog (do not invent cost-per-claim): first-pass / invalid −40%; throughput (280K lines, 11→5); ~$50K lower cost-to-process from paper reduction; $80M program where you refused cheap automation. Same family of units: error rate, volume, cost-to-serve.")

h2(doc, "Canada / BC — gaps and parallels")
body(doc, "Script if they raise US insurance: “I do not have US commercial RCM. I built detect-and-validate on a payer-adjacent system. Candid’s delta is auto-correct and a closed learning loop. That is why I am here.” Then one parallel, then stop.")

h3(doc, "Real gaps (say them cleanly)")
bullet(doc, "No production CPT / ICD-10 / HCPCS, no 837P/837I or 835 ERA, no US clearinghouse.")
bullet(doc, "No Medicare / Medicaid / commercial multi-payer book, no prior auth product, no coordination of benefits, no “which payer does this card actually route to.”")
bullet(doc, "WorkSafeBC is one payer with contracted providers — not 1,000 adversarial payers changing rules.")
bullet(doc, "PIS detected invalids; it did not auto-correct. No flywheel from ops touches into new rules. No claim follow-up on stale AR.")

h3(doc, "Transferable (same failure mode, other side of the payment)")
bullet(doc, "Fee schedules, contracted rates, service codes, eligibility — encoded so bad line items fail at submit (invalid −40%). Candid: same idea, plus write the correction.")
bullet(doc, "Invoice API: orgs only see / bill contracted services. Candid: rosters, credentialing, site-level rules in a knowledge graph.")
bullet(doc, "3.7M-line Pareto, named-cohort launch, PostHog on 322K lines. Candid: data-intensive, do not big-bang payments.")
bullet(doc, "PIS + BCeID GTM/comms. JD: rollouts to customers, ops, leadership.")
bullet(doc, "TSBC dashboards + RFM. JD: reporting for visibility and accountability.")
bullet(doc, "Referral Radar + $80M human-in-the-loop. Candid agents with a residual 5%.")
bullet(doc, "Your named gaps are the product: no auto-correct, no learning, no follow-up, website fee codes, siloed enterprise data. You can describe the missing system because you operated inside it.")

h2(doc, "Touchless rate vs clean claim vs autonomous RCM")
labeled(doc, "Clean claim rate", "Share of claims a clearinghouse accepts. Founders call this a vanity metric — easy to approach 100%, ignores labor and whether the payer paid correctly. (Oak HC/FT; Candid TCR blog.)")
labeled(doc, "Touchless claim rate (TCR)", "Share of claims that go submit → processed → finally adjudicated with no human intervention. Measures revenue quality and cost in one number. Doug Proctor: many customers were at 0% TCR before Candid because a human touched every claim. Legacy systems often cannot measure TCR because they do not log human time.")
labeled(doc, "Autonomous / touchless RCM", "Rules engine + agents prepare, edit, and submit. Corrections happen in software. Humans handle the residual and those touches are supposed to become tomorrow’s rules. Not “no humans in the company.”")
note(doc, "If they say “we’re already at 95%,” do not look disappointed. Say: “Then the PM job is the edge, the flywheel, and enterprise expansion — not reinventing the 95.”")

h2(doc, "If TCR is already 95%+ — where humans still show up")
body(doc, "Investor and company materials cite 95–99% TCR for customers. The remaining slice is the product.")
bullet(doc, "Payer changed a rule or introduced a bug. Batch that used to pay now denies. System should detect and write an upstream rule — until it does, a human is in the loop. (“Adaptive adversaries.”)")
bullet(doc, "Clinical facts that cannot be inferred: missing documentation, unclear medical necessity, attachment required.")
bullet(doc, "Prior authorization never obtained. Software cannot invent an auth number.")
bullet(doc, "Coordination of benefits / which payer (logo ≠ network; Medicaid vs MCO).")
bullet(doc, "Enrollment / out-of-network / wrong site or NPI — unless the graph already has the roster (this is the address example when data is complete).")
bullet(doc, "Appeals and underpayments that need a letter, a record, or a phone call (they are productizing agents here; not 100%).")
bullet(doc, "Customer-authored custom rules that conflict. New codes or time-based rules when the source system still has no field.")
bullet(doc, "Trust / override: biller disagrees with an auto-edit and needs a reason-coded exception.")
note(doc, "Your WorkSafe parallel: 79% always-approved was still a human decision because the 21% and the $80M did not allow rubber-stamping. High automation does not mean zero judgment.")

h2(doc, "How auto-correct works, payers, staying current")
body(doc, "Not magic. A knowledge graph plus a rules engine plus a feedback loop.")
bullet(doc, "Years reverse-engineering billing requirements for 1,000+ US payors — formats, coding preferences, adjudication logic. There is no public logo list of “we support Aetna and UHC only.” Say “1,000+ payors,” not a fake roster.")
bullet(doc, "Graph holds provider rosters, credentialing, fee schedules, charge rules, site/tax-ID context, and observed payer behavior.")
bullet(doc, "At submit: validate and auto-correct fields (routing, provider/location, coding, modifiers, eligibility-related data) so the claim is right before it leaves.")
bullet(doc, "Flywheel: every claim and every human minute is logged. ML suggests new rules from repeated manual work. A denial cluster on a claim family that used to pay triggers a new upstream rule.")
bullet(doc, "Safety: thousands of concurrent rules need versioning, observability, and automated tests so a new rule does not break the rest (“spaghetti”). Palantir-shaped ops, not a spreadsheet of if-then.")
bullet(doc, "Customers can author custom rules. Training exists. That is also how enterprise complexity enters the product.")
bullet(doc, "Freshness: payers change annually/quarterly and ship bugs. The system is designed for continuous evolution, not a yearly fee-schedule PDF upload. That is exactly the gap you named: WorkSafe could not push updated fee schedules/codes into the product; providers hunted the website.")

h2(doc, "How Candid makes money")
bullet(doc, "Sales-led RCM platform, not a self-serve API price list.")
bullet(doc, "Out-of-Pocket (May 2022, sponsored; author is an investor): pricing starts around 3% of collections if the provider codes themselves; about 3% + $6 per claim if Candid does medical coding. Adjustments for state, complexity, volume. Implementation can be extra.")
bullet(doc, "AI Health Index (2026) grades commercial transparency low — Candid does not publish a current rate card. Treat 3% / $6 as directional, not something to quote as official 2026 list price. Interview line: “You make money as a share of collections, so your incentive is the customer getting paid. I would confirm how enterprise and coding are packaged today.”")
bullet(doc, "Why it works with NRR: as a customer’s claim volume grows, Candid’s revenue grows without a new logo. One regrettable churn in seven years is the retention story they tell (founders, Oak).")
bullet(doc, "Series D capital is for taking that motion upmarket and selling automation into legacy stacks — not for reinventing the take-rate.")

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "Part 3 — Your experience → this role")

h2(doc, "WorkSafe / TSBC vs Candid")

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = DARK
        tcPr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EDE9FE")
        shading.set(qn("w:val"), "clear")
        tcPr.append(shading)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            r.font.color.rgb = BODY
    body(doc, "", before=2, after=2)


add_table(
    doc,
    ["Gap you saw at WorkSafeBC", "What you shipped", "What Candid does / the job"],
    [
        [
            "PIS finds errors; does not auto-correct",
            "In-product validation vs fee schedules, rates, service codes. Invalid submissions −40%",
            "Validate and write the correction before submit. Detect was your layer; correct is theirs",
        ],
        [
            "Ops intervenes every time; no learning",
            "Humans stay on $80M extensions; no rule flywheel from those touches",
            "Log human minutes → ML suggests rules → TCR goes up. Closed loop",
        ],
        [
            "No stale-claim follow-up",
            "Submit-side product; AR follow-up was not the system",
            "End-to-end: submit, chase, denial, appeal. Autonomous follow-up is in the category",
        ],
        [
            "Fee schedules / codes on a website",
            "Could not push updates into the billing product",
            "Fee schedules live in the knowledge graph and drive edits automatically",
        ],
        [
            "CBI-class orgs: siloed, location login; no enterprise view; data does not drive action",
            "Portal + org contracts; fragmented for multi-site",
            "Rosters, sites, tax IDs in one graph; dashboards for leadership; address overwrite example",
        ],
        [
            "Dashboards / launch instrumentation",
            "TSBC churn/RFM; PostHog on 322K lines/year",
            "JD: rigorous reporting for visibility and accountability",
        ],
        [
            "Complex change / GTM",
            "PIS named-cohort launch; BCeID/CIAM 5,600 / 0 downtime",
            "JD: rollouts to customers, ops, leadership",
        ],
        [
            "Agent on claim documents",
            "Referral Radar; human retains decision",
            "Agents on eligibility, coding, appeals on top of the graph",
        ],
    ],
)

h2(doc, "One-liners if they poke a row")
labeled(doc, "Dashboards", "I have shipped both: operational data products (TSBC) and in-product funnels (PostHog). I will not launch a billing change I cannot see.")
labeled(doc, "GTM / comms", "PIS and BCeID were change-management products as much as software. Payments path, named cohorts, ops in the room.")
labeled(doc, "Rules", "I encoded payment and eligibility rules. Candid also applies the edit. I am applying for the second half of a system I already understand.")
labeled(doc, "Same metrics", "First-pass yield, lower processing cost, providers back on care. I say cost-to-process and invalid rate, not a fabricated cost-per-claim.")

h2(doc, "Questions this briefing should make you able to ask them")
bullet(doc, "Of the residual 5%, what share is missing clinical data vs payer-rule drift vs prior auth?")
bullet(doc, "When a human corrects a claim today, how fast does that become a suggested rule — and who accepts it, PM or ops?")
bullet(doc, "For multi-site customers, is enterprise visibility a reporting problem or a roster/graph problem?")
bullet(doc, "How do you package pricing now for a group that already has billers vs a digital-health API customer?")

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "Part 4 — What I uniquely bring")
body(doc, "Most PM candidates will talk rules engines and touchless rate. Your edge is having watched providers fail at the interaction layer — who actually uses the tool, which channel they live in, what they cannot see, and why a better system still gets used like the old one. That is payer-adjacent domain observation plus a UX and behavior lens, not “I like design.”")

h2(doc, "The wedge")
body(doc, "I sat on the payer-adjacent side of the same payment. I saw office admins, not clinicians, submit. I saw EMR and Blueprint — not our portal — as the system of record. I saw a high approval rate that still required human judgment. And I saw a launch that did not move because an icon was invisible. Candid’s 95% TCR still dies on those edges. That is the work I already do.")

h2(doc, "Eight observations (say the last line)")

h3(doc, "1. The user of the billing tool is often not the clinician")
labeled(doc, "Observation", "The PXT app was used by office admins. Satisfaction surveys were filled out by physiotherapists — scores mixed contract terms with tool quality.")
labeled(doc, "Proof", "Treatment-extension discovery. Design for the person who submits, not the person on the contract.")
labeled(doc, "Say this", "Candid’s daily user is the biller, not the MD. I research and design for whoever actually hits submit.")

h3(doc, "2. Channel beats UI")
labeled(doc, "Observation", "GPs will not leave OSCAR. Hearing-aid orgs bill through Blueprint. A nicer portal does not unlock their paper.")
labeled(doc, "Proof", "OSCAR 0-to-1: 15 physician interviews, redirected from portal work to EMR-native submit. PIS: one value prop failed hearing aid.")
labeled(doc, "Say this", "I will prioritize integrations and rosters over homepage polish. If the work does not live in your UI, your UI is not the product.")

h3(doc, "3. Discoverability is the product")
labeled(doc, "Observation", "We shipped 11→5 and multi-claim. First month of the target cohort was flat. Users did not see duplicate-line and re-typed like v1.")
labeled(doc, "Proof", "Failed experiment: launch is the start of learning. PostHog, then fix the affordance.")
labeled(doc, "Say this", "An auto-edit that staff override or never notice is the same failure as an invisible icon. I watch real submit behavior after ship.")

h3(doc, "4. Segment GTM, not just the roadmap")
labeled(doc, "Observation", "25 of 3,795 orgs were 50% of remaining paper. We segmented the build, then used one launch message.")
labeled(doc, "Proof", "Hearing aid still had to leave the EMR. Segmentation without messaging is half a strategy.")
labeled(doc, "Say this", "Digital-health API customers and multi-site MSOs are not one rollout. I will not use one comms plan for both.")

h3(doc, "5. High automation is not “automate the rest”")
labeled(doc, "Observation", "79% of 139K treatment extensions were always approved. Senior voices wanted auto-approve. I said no — $80M physio spend.")
labeled(doc, "Proof", "Product designed to strengthen case-manager review, not bypass it.")
labeled(doc, "Say this", "A 95% TCR means the remaining 5% is judgment — prior auth, COB, medical necessity. I have already made that call.")

h3(doc, "6. Invisible decisions create phone volume")
labeled(doc, "Observation", "Providers called because they could not see the case-manager decision or the rationale.")
labeled(doc, "Proof", "We prioritized request tracking and CMS decision rationale over new request types.")
labeled(doc, "Say this", "If Candid blocks or auto-corrects a claim and the biller cannot see why, support becomes the product. Show the reason code.")

h3(doc, "7. Do not add fields the expert already knows")
labeled(doc, "Observation", "Stakeholders wanted a Form 6 body map on extension requests. Physios already know the anatomy. They wanted text on treatment progress.")
labeled(doc, "Proof", "Killed the body map. Lower friction, less CMS translation scope.")
labeled(doc, "Say this", "I will not add chart fields that do not move first-pass. Extra chrome is a tax on the person submitting.")

h3(doc, "8. Identifiers are a UX bug with money attached")
labeled(doc, "Observation", "Manual payee-number entry caused a production incident — wrong payee submitted.")
labeled(doc, "Proof", "Error validation shipped; pre-populate needed auth. Same class of bug as wrong site, NPI, or service address.")
labeled(doc, "Say this", "Candid’s address-overwrite example is this problem at payer scale. I treat identifiers as product, not admin trivia.")

h2(doc, "How I work (if they ask what the lens is)")
bullet(doc, "Watch the real submitter, not the contracted clinician.")
bullet(doc, "Instrument the launch. Features that cannot be found do not exist.")
bullet(doc, "Kill stakeholder chrome that does not change first-pass or cost-to-collect.")
bullet(doc, "Keep humans on expensive judgment. Turn repeated keystrokes into rules.")
note(doc, "Optional extras if they go deep: more lines with a slow grid still feels broken; an AI prototype is not a Kendo-grid estimate; Form 8/11 was frozen in a 2001 billing spec — clinical quality is a billing constraint, not a content project.")

# ══════════════════════════════════════════════════════════════════════════════
h1(doc, "Sources")
bullet(doc, "Oak HC/FT, “How Candid Is Rebuilding the Infrastructure of Medical Billing” (Jul 22, 2026) — TCR, 0%→95%+, address overwrite, time-based coding, adaptive adversaries, one regrettable churn. https://www.oakhcft.com/blog-post/how-candid-is-rebuilding-the-infrastructure-of-medical-billing")
bullet(doc, "Sixth Street, “Why We Invested in Candid Health” — 1,000+ payors, knowledge graph, 95–99% touchless, <20 days to pay, ~40% less manual cost-to-collect. https://sixthstreet.com/blog-post/why-we-invested-in-candid-health-powering-modern-infrastructure-for-a-more-efficient-healthcare-system/")
bullet(doc, "Fortune, Jul 22, 2026 — Series D $120M, ~$219M+ raised, Talkiatry / Nourish proof points, word-of-mouth growth.")
bullet(doc, "Candid, “Why Touchless Claim Rate is the New RCM Standard” — TCR vs clean claim. https://candidhealth.com/blog/beyond-clean-claim-rates-why-touchless-claim-rate-is-the-new-rcm-standard")
bullet(doc, "Out-of-Pocket, “How Healthcare Payments Work with Candid Health” (May 2022, sponsored/investor) — claim path, pains, ~3% of collections / +$6 coding. Confirm packaging. https://www.outofpocket.health/p/how-healthcare-payments-work-with-candid-health")
bullet(doc, "AI Health Index vendor page — API-first, custom rules, commercial transparency gap (no public rate card).")
bullet(doc, "YC / Ashby PM JD — human-centered platform, macro-structure, unit economics, dashboards, rollouts, 4-day office, $150–195k.")
bullet(doc, "Your Candid resume + application: invalid −40%, 3.7M lines, 280K, paper 56%→19%, ~$50K, $580K, 322K lines, $80M / 139K / 79%, ProductCamp 2025.")

out = r"c:\Users\khush\Desktop\kdodani-portfolio\Candid_Health_Interview_Research_Briefing.docx"
doc.save(out)
print(f"Saved: {out}")
