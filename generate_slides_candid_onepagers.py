"""Candid Functional — case-study one-pagers + data lifecycle. One grid, one design system."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Design system
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_BG = RGBColor(0xF5, 0xF3, 0xFF)
GREEN = RGBColor(0x04, 0x78, 0x57)
GREEN_BG = RGBColor(0xEC, 0xFD, 0xF5)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG = RGBColor(0xFE, 0xF2, 0xF2)
GRAY = RGBColor(0xF9, 0xFA, 0xFB)
LINE = RGBColor(0xE5, 0xE7, 0xEB)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MX = Inches(0.40)
GAP = Inches(0.18)
CONTENT_W = SLIDE_W - MX * 2  # 12.533

FONT = "Calibri"


def _run(p, text, size=12, bold=False, color=INK, italic=False, spc=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    if spc is not None:
        rPr = r._r.get_or_add_rPr()
        rPr.set("spc", str(spc))
    return r


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _shape(slide, x, y, w, h, fill=WHITE, line_color=LINE, weight=1.0, roundness=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    _fill(shp, fill)
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(weight)
    try:
        shp.adjustments[0] = roundness
    except Exception:
        pass
    return shp


def _tb(slide, x, y, w, h):
    b = slide.shapes.add_textbox(x, y, w, h)
    tf = b.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return tf


def caption(slide, x, y, w, h, text, size=12, bold=False, color=INK, italic=False, align=PP_ALIGN.LEFT, spc=None):
    tf = _tb(slide, x, y, w, h)
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, bold=bold, color=color, italic=italic, spc=spc)
    return tf


def label(slide, x, y, w, text, color=MUTED):
    caption(slide, x, y, w, Inches(0.20), text.upper(), size=10, bold=True, color=color, spc=180)


def node(slide, x, y, w, h, text, kind="default"):
    styles = {
        "default": (GRAY, LINE, INK),
        "fail": (RED_BG, RED_BG, RED),
        "intervene": (GREEN_BG, GREEN_BG, GREEN),
        "mute": (WHITE, LINE, MUTED),
    }
    fill, border, color = styles[kind]
    line = LINE if kind in ("default", "mute") else border
    _shape(slide, x, y, w, h, fill=fill, line_color=line, weight=1.0)
    tf = _tb(slide, x + Inches(0.06), y + Inches(0.06), w - Inches(0.12), h - Inches(0.12))
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _run(tf.paragraphs[0], text, size=11, bold=True, color=color)


def rarrow(slide, x, y, color=MUTED):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, Inches(0.20), Inches(0.14))
    _fill(shp, color)
    shp.line.fill.background()


def flow_row(slide, nodes, y, kinds=None):
    n = len(nodes)
    kinds = kinds or ["default"] * n
    arrow_w = Inches(0.28)
    node_w = (CONTENT_W - arrow_w * (n - 1)) / n
    h = Inches(0.44)
    x = MX
    for i, text in enumerate(nodes):
        node(slide, x, y, node_w, h, text, kind=kinds[i])
        if i < n - 1:
            ac = GREEN if kinds[i] == "intervene" or kinds[i + 1] == "intervene" else (
                RED if kinds[i] == "fail" or kinds[i + 1] == "fail" else MUTED
            )
            rarrow(slide, x + node_w + Inches(0.04), y + Inches(0.15), color=ac)
        x += node_w + arrow_w


def panel(slide, x, y, w, h, title, lines, fill=GRAY, title_color=INK):
    _shape(slide, x, y, w, h, fill=fill, line_color=LINE, weight=1.0)
    pad = Inches(0.20)
    tf = _tb(slide, x + pad, y + Inches(0.12), w - pad * 2, h - Inches(0.22))
    p = tf.paragraphs[0]
    _run(p, title.upper(), size=10, bold=True, color=title_color, spc=180)
    for line in lines:
        bp = tf.add_paragraph()
        bp.space_before = Pt(4)
        _run(bp, line, size=12, color=INK)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _fill(bg, WHITE)
    bg.line.fill.background()
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.045))
    _fill(bar, PURPLE)
    bar.line.fill.background()
    return s


def header(slide, kicker, title, user, business, trigger, ideation,
           l1="User", l2="Business", l3="Trigger"):
    caption(slide, MX, Inches(0.10), CONTENT_W, Inches(0.18),
            kicker.upper(), size=10, bold=True, color=PURPLE, spc=200)
    caption(slide, MX, Inches(0.26), CONTENT_W, Inches(0.34),
            title, size=22, bold=True, color=INK)
    tf = _tb(slide, MX, Inches(0.58), CONTENT_W, Inches(0.18))
    p = tf.paragraphs[0]
    _run(p, f"{l1}  ", size=10, bold=True, color=MUTED)
    _run(p, user + "     ·     ", size=10, color=MUTED)
    _run(p, f"{l2}  ", size=10, bold=True, color=MUTED)
    _run(p, business + "     ·     ", size=10, color=MUTED)
    _run(p, f"{l3}  ", size=10, bold=True, color=MUTED)
    _run(p, trigger, size=10, color=MUTED)
    caption(slide, MX, Inches(0.74), CONTENT_W, Inches(0.18),
            ideation, size=11, italic=True, color=MUTED)


def insight(slide, y, text):
    h = Inches(0.42)
    _shape(slide, MX, y, CONTENT_W, h, fill=PURPLE_BG, line_color=None)
    caption(slide, MX + Inches(0.22), y + Inches(0.08), CONTENT_W - Inches(0.44), Inches(0.28),
            text, size=14, bold=True, color=PURPLE, align=PP_ALIGN.LEFT)


def outcome_bar(slide, y, metrics, why, left_title="OUTCOME"):
    h = Inches(0.78)
    left_w = Inches(8.55)
    right_w = CONTENT_W - left_w - GAP
    _shape(slide, MX, y, left_w, h, fill=GREEN_BG, line_color=None)
    caption(slide, MX + Inches(0.22), y + Inches(0.10), left_w - Inches(0.40), Inches(0.18),
            left_title, size=10, bold=True, color=GREEN, spc=180)
    caption(slide, MX + Inches(0.22), y + Inches(0.32), left_w - Inches(0.40), Inches(0.36),
            metrics, size=13, bold=True, color=INK)

    rx = MX + left_w + GAP
    _shape(slide, rx, y, right_w, h, fill=GRAY, line_color=LINE)
    caption(slide, rx + Inches(0.18), y + Inches(0.10), right_w - Inches(0.32), Inches(0.18),
            "WHY THIS MATTERS", size=10, bold=True, color=MUTED, spc=180)
    caption(slide, rx + Inches(0.18), y + Inches(0.30), right_w - Inches(0.32), Inches(0.40),
            why, size=12, bold=False, color=INK)


def footer(slide, n, scope):
    caption(
        slide, MX, Inches(7.18), Inches(8.6), Inches(0.22),
        f"Healthcare systems fail upstream. I fix input quality instead of scaling downstream ops.    ·    {n}/6",
        size=10, color=MUTED,
    )
    caption(
        slide, Inches(9.10), Inches(7.18), Inches(3.80), Inches(0.22),
        f"Scope  {scope}",
        size=10, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def two_col(slide, y, h, left_title, left_lines, right_title, right_lines):
    col_w = (CONTENT_W - GAP) / 2
    panel(slide, MX, y, col_w, h, left_title, left_lines, fill=RED_BG, title_color=RED)
    panel(slide, MX + col_w + GAP, y, col_w, h, right_title, right_lines, fill=GREEN_BG, title_color=GREEN)


def three_col(slide, y, h, cols):
    col_w = (CONTENT_W - GAP * 2) / 3
    x = MX
    for title, lines, tcolor in cols:
        panel(slide, x, y, col_w, h, title, lines, fill=GRAY, title_color=tcolor)
        x += col_w + GAP


def data_grid(slide, y, row_h, rows):
    """rows: (label, [5 cell strings], label_color). Column titles sit above the cells."""
    headers = ["Discover", "Prioritize", "Decide", "Launch / learn", "Outcome"]
    label_w = Inches(0.88)
    gap = Inches(0.08)
    n = 5
    cell_w = (CONTENT_W - label_w - gap * n) / n

    x = MX + label_w + gap
    for title in headers:
        caption(slide, x, y, cell_w, Inches(0.18), title.upper(),
                size=10, bold=True, color=MUTED, spc=180, align=PP_ALIGN.CENTER)
        x += cell_w + gap

    y = y + Inches(0.20)
    for row_label, cells, lcolor in rows:
        _shape(slide, MX, y, label_w, row_h, fill=GRAY, line_color=LINE)
        caption(
            slide, MX, y + Inches((row_h.inches - 0.20) / 2), label_w, Inches(0.20),
            row_label, size=12, bold=True, color=lcolor, align=PP_ALIGN.CENTER,
        )
        x = MX + label_w + gap
        for text in cells:
            _shape(slide, x, y, cell_w, row_h, fill=WHITE, line_color=LINE)
            caption(
                slide, x + Inches(0.08), y + Inches(0.08),
                cell_w - Inches(0.16), row_h - Inches(0.14),
                text, size=11, color=INK,
            )
            x += cell_w + gap
        y = y + row_h + gap


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Shared vertical rhythm
    y_before_lbl, y_before = Inches(1.02), Inches(1.22)
    y_after_lbl, y_after = Inches(1.74), Inches(1.94)
    y_insight = Inches(2.50)
    y_grid1, h_grid1 = Inches(3.04), Inches(1.40)
    y_grid2, h_grid2 = Inches(4.58), Inches(1.34)
    y_outcome = Inches(6.06)

    # ── SLIDE 1 PIS ───────────────────────────────────────────────────────────
    s = blank(prs)
    header(s, "1  ·  first-pass yield",
           "PIS — Product blocked the segment that drove paper",
           user="Delayed pay / physio pre-auth after denials",
           business="$500K ops · 60% denials",
           trigger="50% still paper; paying to transcribe",
           ideation="Ops assumed behavior. The product was built for occasional 20–30 line claims.")

    label(s, MX, y_before_lbl, Inches(2.2), "Before", RED)
    flow_row(s, ["Paper / email", "Ops / ELAN encode", "CMS deny"],
             y_before, kinds=["fail", "fail", "fail"])

    label(s, MX, y_after_lbl, Inches(2.2), "After", GREEN)
    flow_row(s, ["High-volume PIS", "★ Validation", "CMS"],
             y_after, kinds=["default", "intervene", "mute"])

    insight(s, y_insight,
            "Not a training problem. Architecture blocked the 25 providers who were 50% of paper.")

    two_col(
        s, y_grid1, h_grid1,
        "Problem",
        [
            "$500K ops cost; 60% denials",
            "Providers cleaning up after deny",
            "Delayed pay / physio pre-auth",
            "Paying a vendor to transcribe a “digital” channel",
        ],
        "Discovery",
        [
            "Ops: marketing / pay-to-use / tax paper",
            "Product: 1-claim / 32-line vs ~1000 multi-claim",
            "25 orgs = 50% of paper, similar workflow",
            "Segment + workflow, not features for the average user",
        ],
    )

    three_col(s, y_grid2, h_grid2, [
        ("Shipped", [
            "Fewer screens + multi-claim",
            "Fee codes / programs + contract rules at submit",
            "1000–2000 line submissions",
            "Async CMS + retry",
        ], GREEN),
        ("Launch / failure", [
            "10-name dark launch: too slow to submit",
            "Optimized capacity, not responsiveness",
            "Dual-run v1 live → soft → hard",
        ], RED),
        ("Tradeoffs", [
            "25 similar-workflow first vs all providers",
            "Dark/dual vs big bang → dual-run",
            "Portal vs EMR; sync vs async → async",
        ], INK),
    ])

    outcome_bar(
        s, y_outcome,
        "Invalid −40% (user)    ·    Adoption +60% (behavior)    ·    $200K ops saved (business)",
        "Denial and ops cost start at submit, not at marketing.",
    )
    caption(
        s, MX, Inches(6.88), CONTENT_W, Inches(0.22),
        "Missed: no auto-correct / rules not reconfigurable    ·    no sandbox vs live CMS deny    ·    bulk multi-claim (didn’t push line-level pay)",
        size=10, color=MUTED,
    )
    footer(s, 1, "25-org Pareto · legacy Kendo / CMS")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK
Business: $500K ops, 60% denials. Providers unpaid; physio delayed on pre-auth.
Why paper if we have a portal? 50% still paper. Paying ELAN to transcribe.
Ops assumed behavior — more marketing, pay-to-use, tax paper. That bleeds into ops cost too.
Discovery: product fit 20–30 single-claim lines, not 1000-line multi-claim. 25 orgs = 50% of paper, similar workflow. Unlock a segment; don’t spray features on the average user.
Shipped: fewer screens, multi-claim, fee codes/programs, contract rules at submit, 1000–2000 lines, async + retry.
Failure: 10-name dark launch — they never submitted. Too slow, poor UX. I optimized capacity, not responsiveness. Scalability and performance are one problem.
Then dual-run so existing users stayed live; soft → hard.
Outcomes: invalid −40%, adoption +60%, $200K ops saved. Do not quote a post-denial % — we don’t have one.
Missed: no auto-correct, rules not reconfigurable, no sandbox vs live CMS deny, bulk multi-claim (didn’t push line-level pay). Mandatory docs as adoption blocker. That’s the Candid conversation.
Strengths (if asked): Pareto 25; eng/architects/rules-engine PM/CMS/CMs/marketing; InfoX was born from this (API-first).

PROBE
Capacity vs responsiveness. Why dual-run. Why portal before EMR. What you’d sandbox vs live deny.
"""

    # ── SLIDE 2 INFOX ─────────────────────────────────────────────────────────
    s = blank(prs)
    header(s, "2  ·  integrations",
           "InfoX — Clinical data trapped in a billing spec",
           user="GPs documenting in the EMR",
           business="52K faxes/yr · ~$250K re-key",
           trigger="Notes truncated by 2001 HIBC",
           ideation="Assumed a better form or a portal. GPs will not leave the EMR.")

    label(s, MX, y_before_lbl, Inches(2.2), "Before", RED)
    flow_row(s, ["Doctor / EMR", "PDF / eForm", "Fax", "ELAN re-key", "CMS"],
             y_before, kinds=["default", "default", "fail", "mute", "mute"])

    label(s, MX, y_after_lbl, Inches(2.2), "After", GREEN)
    flow_row(s, ["EMR eForm", "★ API / APIM", "CMS (clinical)", "Teleplan (bill)"],
             y_after, kinds=["default", "intervene", "default", "mute"])

    insight(s, y_insight, "Billing was the wrong abstraction for clinical data.")

    two_col(
        s, y_grid1, h_grid1,
        "Problem",
        [
            "Form 8/11 locked in 2001 Teleplan/HIBC",
            "Form frozen since 2015; notes truncated",
            "Digital in EMR ≠ digital in CMS",
            "52K faxes; paying ELAN to re-key",
        ],
        "Discovery",
        [
            "15 GP interviews: they will not leave the EMR",
            "Advisors: clinical text cut off by billing limits",
            "InnoFest “Death to PDFs” — didn’t win",
            "$80K Innovation Council POC anyway",
        ],
    )

    three_col(s, y_grid2, h_grid2, [
        ("Shipped", [
            "Split clinical off Teleplan/HIBC",
            "OSCAR eForm → APIM → CMS JSON",
            "APIM facade reusable for any EMR",
            "Billing stays on Teleplan",
        ], GREEN),
        ("Launch / failure", [
            "Form/portal instinct died in discovery",
            "OSCAR POC validated 27 Mar 2025 — F8 not live",
            "Vendor talks sequenced after the POC",
        ], RED),
        ("Tradeoffs", [
            "Patch form vs rebuild pipe → rebuild",
            "FHIR now vs JSON now → JSON",
            "Top-5 vs OSCAR beachhead → OSCAR",
        ], INK),
    ])

    outcome_bar(
        s, y_outcome,
        "GPs stay in EMR (user)    ·    $80K POC · 27 Mar 2025 (business)    ·    Not a fax-% claim",
        "Clean inputs beat downstream automation. Distribution lives in the EMR.",
    )
    caption(
        s, MX, Inches(6.88), CONTENT_W, Inches(0.22),
        "Missed: vendor talks parallel with discovery    ·    live Form 8 not in production    ·    FHIR still the upgrade path",
        size=10, color=MUTED,
    )
    footer(s, 2, "EMR vendor constraints · JSON not FHIR")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK
Business: 52K faxes/year, ~$250K re-key. Clinical notes truncated by a 2001 billing spec. Form frozen since 2015.
Assumption: better form, portal, or pay EMRs. Discovery killed it — 15 GPs will not leave the EMR. Advisors: text cut off by HIBC limits. Digital in EMR ≠ digital in CMS.
Reframe: billing was the wrong abstraction for clinical data. Split the pipe; don’t patch the 2015 form.
Shipped: OSCAR eForm → APIM → CMS JSON. Teleplan stays billing. APIM is reusable.
Launch: InnoFest “Death to PDFs” didn’t win; $80K POC anyway. Validated 27 Mar 2025. Form 8 not live. Vendor talks ran after the POC — that was the miss.
GTM (if asked): EMRs integrate to keep doctors. Pull demand, don’t pay for the pipe. Path to ~70%; signals from QHR, Plexia, Medinet, Intrahealth. 50 vendors in outreach.
Outcome: GPs stay in EMR. $80K POC. Not a fax-% claim.
Missed: vendor talks in parallel with discovery; live F8; FHIR later.

PROBE
Why JSON not FHIR. Why OSCAR not top-5. How you create pull as a small fish.
"""

    # ── SLIDE 3 RADAR ─────────────────────────────────────────────────────────
    s = blank(prs)
    header(s, "3  ·  residual 5%",
           "Referral Radar — Finding was the bottleneck, not judging",
           user="CMs on referral calls",
           business="90+ min SME hunt",
           trigger="Signal buried in 20–75 docs",
           ideation="Assumed a ranked recommendation was the product. Explainability was.")

    label(s, MX, y_before_lbl, Inches(2.2), "Before", RED)
    flow_row(s, ["Claim docs", "Hunt 20–75", "CM decides", "Referral"],
             y_before, kinds=["default", "fail", "default", "mute"])

    label(s, MX, y_after_lbl, Inches(2.2), "After", GREEN)
    flow_row(s, ["Claim docs", "★ Retrieval", "Evidence", "CM decides", "Referral"],
             y_after, kinds=["default", "intervene", "default", "default", "mute"])

    insight(s, y_insight, "This is retrieval + evidence, not automation.")

    two_col(
        s, y_grid1, h_grid1,
        "Problem",
        [
            "20–75 docs per claim; 90+ min finding",
            "Bottleneck is finding, not judging",
            "No systematic way to connect symptoms, barriers, eligibility",
            "CMs will not act without evidence",
        ],
        "Discovery",
        [
            "Trust requires an evidence trace, not a score",
            "Production PHI unavailable in the sprint",
            "Synthetic data unblocked DS",
            "36 InnoFest ideas; 2-week prototype",
        ],
    )

    three_col(s, y_grid2, h_grid2, [
        ("Shipped", [
            "Retrieval over the document pile",
            "Evidence + mind map to the CM",
            "Blocked auto-decision by design",
            "CM-validated synthetic data (no production PHI)",
        ], GREEN),
        ("Launch / failure", [
            "2-week InnoFest prototype — not production",
            "Assumed accuracy would land; stakeholders used the map",
            "No confidence / override patterns in the sprint",
        ], RED),
        ("Tradeoffs", [
            "Auto-refer vs human-in-loop → human",
            "Ranked list vs mind map → mind map",
            "Wait on PHI vs synthetic → synthetic",
        ], INK),
    ])

    outcome_bar(
        s, y_outcome,
        "Evidence to CM, not a score (user)    ·    Prototype · 2nd of 36 (signal)    ·    No minutes-saved claim",
        "Trust is the constraint, not model accuracy. Residual 5% stays human.",
    )
    caption(
        s, MX, Inches(6.88), CONTENT_W, Inches(0.22),
        "Missed: confidence / override earlier    ·    no production validation    ·    CM touches → rules not shipped",
        size=10, color=MUTED,
    )
    footer(s, 3, "No production PHI · 2-week sprint")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK
Business: CMs hunting 20–75 docs, 90+ min SME. Signal buried. ~70 programs if they ask.
Assumption: a ranked recommendation (accuracy) was the product. Discovery: finding, not judging, is the bottleneck. Trust needs an evidence trace. No production PHI in the sprint.
Reframe: retrieval + evidence, not automation. Human-in-loop — same residual call as $80M extensions.
Shipped: retrieval layer, mind map + evidence, blocked auto-decision, CM-validated synthetic data.
Failure: 2-week prototype, not production. Assumed accuracy would land; stakeholders used the map. No override patterns in the sprint.
Outcome: 2nd of 36. Evidence to the CM, not a score. No minutes-saved claim.
Missed: confidence/override earlier; production validation; CM touches → rules.

PROBE
Why not auto-route obvious cases. Synthetic edge misses. Who accepts a rule from a CM touch.
"""

    # ── SLIDE 4 VISION ────────────────────────────────────────────────────────
    s = blank(prs)
    header(
        s, "4  ·  vision  ·  HoP / Product Lead",
        "After failure → change the system that produces claims",
        user="AI denial arms race",
        business="Residual tail at high touchless",
        trigger="Claims still born wrong",
        ideation="Market and law converge on assist, don’t replace.",
        l1="Industry", l2="Constraint", l3="Trigger",
    )

    label(s, MX, y_before_lbl, Inches(3.4), "Candid today", RED)
    flow_row(
        s, ["Fail", "Agent executes", "Rule", "Next claim"],
        y_before, kinds=["fail", "default", "default", "mute"],
    )

    label(s, MX, y_after_lbl, Inches(3.4), "Next layer", GREEN)
    flow_row(
        s, ["Pattern emerging", "Infer cause", "Change default", "Claim never wrong"],
        y_after, kinds=["default", "intervene", "intervene", "mute"],
    )

    insight(
        s, y_insight,
        "Candid learns from resolved errors. Next: learn from emerging patterns before they scale.",
    )

    two_col(
        s, y_grid1, h_grid1,
        "Where the industry breaks",
        [
            "Payers auto-deny on missing data / coding in seconds",
            "Providers racing to pre-submit prediction",
            "Regulatory backlash on AI denials → HITL, not full autonomy",
            "Residual cases get harder as touchless rises",
        ],
        "What I’d add",
        [
            "Decision agents that change defaults / workflows, not just fix claims",
            "Infer provider-level miscode as intake/training issue",
            "Payer × CPT drift across the customer base",
            "Integrate EMR-embedded billing; don’t replace it",
        ],
    )

    three_col(s, y_grid2, h_grid2, [
        ("On Candid’s roadmap", [
            "Execution agents (portal, policy, hold, stale)",
            "Push notify + cited Q&A",
            "Error → rule flywheel",
            "HITL on the residual",
        ], INK),
        ("Decision", [
            "Extend execution into proactive intelligence",
            "Assist, don’t replace (regulation agrees)",
            "Iterations: fail-loop → pattern-loop → config",
        ], PURPLE),
        ("Tradeoffs", [
            "Full autonomy vs HITL → HITL (law + trust)",
            "React to denials vs preempt at intake → preempt",
            "All agents vs one payer × CPT → beachhead",
        ], INK),
    ])

    outcome_bar(
        s, y_outcome,
        "Execution you have    ·    Decision layer I’d add    ·    GTM: EMR-embedded billing",
        "Touchless flywheel + cost-to-serve on the residual.",
        left_title="BET",
    )
    footer(s, 4, "Thesis, not a spec · HoP / Product Lead")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK (HoP / Product Lead)
Industry: both sides in an AI arms race — payers auto-deny in seconds; providers racing to pre-submit prediction. Law is pushing HITL, not full autonomy. That’s already Candid’s design principle: assist, don’t replace.

Point at the flow: you compressed the loop AFTER a claim fails — call payer, policy, denial reason, stale claim. I’d pull that loop BEFORE the claim is born.

Sharp line: Candid learns from resolved errors. The next layer learns from emerging patterns. Candid optimizes claim correctness. Next layer optimizes the system behavior that produces claims. (I call this JEV in my notes — not a competing product.)

Scenario: provider miscodes CPT X from a bad intake default.
Candid: detect denial → fix claim → add rule.
Next layer: this provider consistently miscodes → infer workflow/training → suggest UI default, flag provider, change intake — so those claims never get generated incorrectly.

ML: your flywheel already learns from resolved errors across the customer base. Point the same aggregate signal at payer × CPT drift — policy as something that drifts, not something fixed until a human notices a spike.

GTM: EMR-embedded billing is high-volume, high-complexity — where revenue and denial risk sit. You don’t replace those systems; you integrate. Same shape as WorkSafe: 25 orgs, 50% of paper. Solving for them compounds FPY.

Gamification (one line, if they ask): you’re building a biller org for residual HITL. Visible speed/accuracy on that queue is a companion to agents, not a separate bet.

STOP. Ask: what are you already building on the decision layer vs execution?

PROBE
1. At high touchless, is the bottleneck validation or how you handle the long tail?
2. Would you start decision agents on one payer × CPT, or horizontal across agents?
3. How do you currently detect payer policy drift vs waiting on a denial spike?
"""

    # ── SLIDE 5 DATA LIFECYCLE ────────────────────────────────────────────────
    s = blank(prs)
    header(
        s, "5  ·  how I use data",
        "Data changes the product — it does not decorate the deck",
        user="Interviews + volume, not dashboards",
        business="PHI, legacy specs, average-user product",
        trigger="Ops assumed behavior; I used data to kill the wrong product",
        ideation="Same loop on three problems. Volume finds the constraint. Users kill the form, the average user, and auto-refer.",
        l1="Method", l2="Constraint", l3="Trigger",
    )

    label(s, MX, Inches(1.02), Inches(3.4), "Lifecycle", INK)
    flow_row(
        s, ["Discover", "Prioritize", "Decide", "Launch / learn", "Outcome"],
        Inches(1.22),
        kinds=["default", "default", "intervene", "fail", "mute"],
    )

    insight(
        s, Inches(1.80),
        "Volume finds where it breaks. Interviews kill the wrong product. Launch data kills the wrong success metric.",
    )

    data_grid(s, Inches(2.34), Inches(1.08), [
        ("PIS", [
            "3.7M lines · 25 orgs = 50% paper · $500K · 60% denials",
            "25 similar-workflow first · product fit 20–30 vs ~1000 lines",
            "Architecture, not training · 32-line → 1000–2000",
            "10-name dark too slow · capacity ≠ responsiveness · dual-run",
            "Invalid −40% · adoption +60% · $200K · no post-denial %",
        ], PURPLE),
        ("InfoX", [
            "52K faxes · ~$250K re-key · 28M keystrokes · 15 GPs + advisors",
            "Form 8/11 first (highest-volume clinical) · $80K POC",
            "GPs will not leave EMR · split clinical off Teleplan · JSON not FHIR",
            "InnoFest loss · 27 Mar 2025 POC · F8 not live",
            "Channel validated · GPs stay in EMR · not a fax-% claim",
        ], INK),
        ("Radar", [
            "20–75 docs · 90+ min SME (not an ops extract) · finding, not judging",
            "Explainability over a ranked score · 2nd of 36",
            "Evidence trace · human-in-loop · mind map not list",
            "No production PHI · synthetic + CM validate · 2-week prototype",
            "Evidence to CM, not a score · no minutes-saved · not production",
        ], GREEN),
    ])

    outcome_bar(
        s, y_outcome,
        "Same muscle on three stages of a claim: submit (PIS) · clinical intake (InfoX) · residual 5% (Radar)",
        "Candid’s rules flywheel is this loop: volume at submit, user evidence, refuse vanity.",
    )
    caption(
        s, MX, Inches(6.88), CONTENT_W, Inches(0.22),
        "Missed: Radar 20–75 / 90+ not ops-validated    ·    InfoX no vendor-pipeline metric    ·    PIS no sandbox vs live deny",
        size=10, color=MUTED,
    )
    footer(s, 5, "Locked facts only · no invented TCR")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK
The loop: volume finds the constraint. Interviews kill the wrong product. Launch data kills the wrong success metric. Same muscle on submit (PIS), clinical intake (InfoX), residual 5% (Radar).

PIS: 3.7M lines → 25 orgs = 50% paper. Ops assumed behavior; product was 20–30 line. I funded the segment, not the average user. Dark launch of 10 names was too slow — capacity ≠ responsiveness. Outcomes I will quote: invalid −40%, adoption +60%, $200K. I will not quote a post-denial %.

InfoX: 52K faxes, ~$250K re-key, 28M keystrokes. 15 GPs + advisors: they will not leave the EMR. Form 8/11 first. $80K POC. JSON not FHIR. InnoFest loss; validated 27 Mar 2025; Form 8 not live. Not a fax-% claim.

Radar: 20–75 docs, 90+ min from SMEs — not an ops extract. Finding, not judging. No production PHI → synthetic + CM validate. Evidence to the CM, not a score. 2nd of 36. No minutes-saved. Not production.

What I will not claim: fax-% drop, minutes saved, post-denial %, a TCR number I do not own.

PROBE
How do you instrument TCR vs the residual? Who owns the metric when PHI blocks the model? What vanity number would you refuse to put on a slide?
"""

    # ── SLIDE 6 PIS ARCHITECTURE ──────────────────────────────────────────────
    s = blank(prs)
    header(
        s, "6  ·  architecture",
        "PIS — Fail closed at submit, retry after CMS",
        user="Office admin (not the MD)",
        business="Invalid never reaches ops",
        trigger="Portal can succeed while CMS fails",
        ideation="Detect-and-validate. Entitlement stays human. V1 stayed live on purpose.",
    )

    label(s, MX, y_before_lbl, Inches(2.2), "Data", GREEN)
    flow_row(
        s, ["Office admin", "PIS V2 / Kendo", "★ Validation", "PIS.API", "async CMS"],
        y_before,
        kinds=["default", "default", "intervene", "default", "mute"],
    )

    label(s, MX, y_after_lbl, Inches(2.2), "Error", RED)
    caption(
        s, Inches(2.70), y_after_lbl, Inches(10.2), Inches(0.20),
        "CMS fail  →  retry queue  →  eng email for manual submit",
        size=10, color=MUTED, italic=True,
    )
    flow_row(
        s, ["No DOS / ineligible", "★ Fail closed", "never queued"],
        y_after,
        kinds=["fail", "fail", "fail"],
    )

    insight(s, y_insight, "Fail closed at submit. CMS can still fail after the portal succeeds.")

    three_col(s, y_grid1, Inches(2.70), [
        ("No DOS", [
            "Fee codes live in CMS HUB with multiple periods",
            "DOS is the join key",
            "No date → no fee-code search",
        ], RED),
        ("Ineligible", [
            "Invoice API: orgs only see contracted services",
            "Invalid never queued",
            "Cut polish so eligibility did not wreck submit latency",
        ], RED),
        ("Portal OK / CMS fail", [
            "Sync → async to scale",
            "Retry; then email to eng for manual submit",
            "Residual: “is this line for this injury?” stays human, not a regex",
        ], RED),
    ])

    outcome_bar(
        s, y_outcome,
        "Invalid −40% (user)    ·    Fail closed at submit (system)    ·    Residual stays human",
        "Same detect layer as Candid. Missing half: auto-correct + learn from the touch.",
    )
    caption(
        s, MX, Inches(6.88), CONTENT_W, Inches(0.22),
        "Missed: no sandbox vs live CMS deny    ·    rules not reconfigurable    ·    no auto-correct",
        size=10, color=MUTED,
    )
    footer(s, 6, "Detect only · legacy Kendo / CMS")
    s.notes_slide.notes_text_frame.text = """5–7 MIN WALK
Walk the pipe left to right. Office admin, not the MD. V1 stayed live on purpose.

DATA: PIS V2 / Kendo → validation (fee codes, contract, DOS) → PIS.API → async CMS.

ERROR 1: fail closed at submit. No DOS → no fee-code search (CMS HUB periods; DOS is the join key). Ineligible contract / invalid code → never queued. I bring the invalid state; eng brings the constraint. Cut polish so eligibility did not wreck submit latency.

ERROR 2: portal can succeed while CMS fails. Sync → async to scale. Retry queue; then email to eng for manual submit (WOR pattern).

Residual: grid encodes format and contract. “Is this line for this injury?” stays human — not a regex.

Invalid −40% because bad lines never reached ops. Detect only. No auto-correct. Do not claim a measured timezone-denial rate. Dark-launch slowness is capacity vs responsiveness — not a box on this slide.

PROBE
What do you sandbox vs live deny? Where does entitlement stay human on your stack?
"""

    out = r"c:\Users\khush\Desktop\kdodani-portfolio\Candid_Health_Functional_OnePagers.pptx"
    prs.save(out)
    print(f"Saved: {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
