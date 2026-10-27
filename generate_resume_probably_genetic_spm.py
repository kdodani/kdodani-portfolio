from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)

DARK    = RGBColor(0x1c, 0x19, 0x17)
MID     = RGBColor(0x5B, 0x21, 0xB6)
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)
BODY    = RGBColor(0x44, 0x40, 0x3c)
LIGHT   = RGBColor(0x78, 0x71, 0x6c)
BG_HEX  = "FAF9F6"

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=8, space_after=2, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=6, after=1)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=2, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=1, after=2)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("SENIOR PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Product Bets  •  Growth  •  Healthcare Data Platforms  •  AI-Native Tooling"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product owner for healthcare data platforms used by 5,600+ providers — the line that replaces paper, fax, "
    "and side-portal work with in-workflow digital products on treatment pathways. 5 years placing product bets, "
    "killing the wrong ones, and growing adoption against a real market (fax vs digital; portal vs EMR). "
    "AI-native: prototypes and lightweight automations in Cursor, GitHub Copilot, Lovable, and n8n. Master of "
    "Business Analytics (SQL, Python); BBA in Marketing.",
    size=9.5,
    space_before=4,
    space_after=4,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

subheading(doc, "Product Manager", "WorkSafeBC (via independent consulting engagement)", "Apr 2024 – Present")
body(
    doc,
    "Own the product line healthcare organizations use to submit invoices, referrals, and treatment requests — "
    "5,600 healthcare professionals across 500 organizations.",
    size=9, color=LIGHT, space_before=0, space_after=3,
)

sublabel(doc, "Product bets and growth")
bullet(doc,
    "Bet against the long tail: 25 of 3,795 organizations generated 50% of remaining paper (~280K line items). "
    "Rebuilt invoicing for those accounts — 11 screens to 5 — and grew portal share 39% → 55% while paper fell "
    "56% → 19% of line items.")
bullet(doc,
    "Bet on in-workflow EMR over a side portal: physicians will not leave OSCAR to file clinical documentation. "
    "Funded a 0-to-1 path covering ~70% of the provider ecosystem.")
bullet(doc,
    "Grew the treatment-extension product off fax/phone onto digital treatment pathways — an $80M physiotherapy "
    "line where stalled requests stop patient outcomes.")
bullet(doc,
    "Unlocked $1.2M+ to fund the line (invoicing, identity, EMR integrations), each case defended to an "
    "investment committee.")

sublabel(doc, "Tradeoffs")
bullet(doc,
    "Killed auto-approve after 139K extension requests showed 79% always approved — speed would have rubber-stamped "
    "an $80M spend; clinical judgment stayed in the product.")
bullet(doc,
    "Launched by named cohort (dark / soft / hard) so existing digital volume kept moving while the paper segment "
    "activated — growth without a big-bang cutover.")
bullet(doc,
    "Held a hard human-in-the-loop boundary on Referral Radar: AI retrieves clinical data, humans decide — the "
    "tradeoff that made the pilot approvable.")

sublabel(doc, "Insight that changed the bet")
bullet(doc,
    "The 25-account segmentation reframed a stalled adoption program into a growth bet against named organizations, "
    "not a generic change campaign.")
bullet(doc,
    "Instrumented PostHog before launch so growth, drop-off, and activation were visible within days.")
bullet(doc,
    "Support calls down 30% and clinician NPS +25 after validation at submission — data quality as a growth lever, "
    "not a cleanup job.")

sublabel(doc, "AI-native build")
bullet(doc,
    "Built Referral Radar in two weeks — LLM/RAG over 20–75 documents per case, evidence traces, 2nd of 36 at "
    "InnoFest 2025. Unblocked with SME-validated synthetic data when production access was blocked.")
bullet(doc,
    "Prototype in Cursor, GitHub Copilot, and Lovable; automate reporting in n8n rather than wait on a ticket.")
bullet(doc,
    "Migrated 5,600 healthcare professionals across 500 organizations to modern identity on a portal carrying "
    "62,000+ annual referrals — 0 downtime.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Bet on a mobile inspections product after field research — mapped the workflow and tested prototypes before "
    "spending engineering.")
bullet(doc,
    "Treated retention as the growth problem: 11.7% annual churn, RFM campaigns that kept 78% of high-risk customers.")
bullet(doc,
    "Shipped NLP inbox routing into live operations — cut manual triage without a separate build program.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Delivered scoped consulting engagements — customer research, market analysis, and presentations to client "
    "leadership; designed a data-driven acquisition strategy for Open Primaries that cut bounce rate 37% while "
    "doubling subscriptions.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Advised 75+ client websites on analytics-driven growth, increasing organic performance ~150%, and built a "
    "forecasting tool predicting SEO performance with ~85% accuracy.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")

sublabel(doc, "Speaker — ProductCamp BC (2025)", italic=False, size=9.5, color=DARK, space_before=0)
body(doc,
    '"How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery."',
    size=9.5, italic=True, space_before=0)

sublabel(doc, "Mentor — ProductBC", italic=False, size=9.5, color=DARK, space_before=4)
body(doc,
    "Mentor aspiring Product Managers in product discovery, strategy, and career development.",
    size=9.5, space_before=0)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY

skills_row(doc, "Product",
    "Product Bets  •  Tradeoffs  •  Growth  •  Product Line Ownership  •  0-to-1  •  Market Framing")
skills_row(doc, "Customer & Insights",
    "Voice of the Customer  •  Named-Account Targeting  •  Insight Packages  •  Executive Presentations  •  "
    "Change Management")
skills_row(doc, "Analytics",
    "SQL  •  Python  •  Excel  •  PostHog  •  Funnel & Cohort Analysis  •  Segmentation  •  Business Cases")
skills_row(doc, "AI & Automation",
    "LLM / RAG  •  n8n  •  Rapid Prototyping  •  Synthetic Data  •  Human-in-the-Loop Design")
skills_row(doc, "Healthcare Data",
    "Healthcare Data Platforms  •  Insurance Operations  •  Provider Networks  •  EMR/EHR Integrations  •  "
    "Regulated Environments")
skills_row(doc, "Tools",
    "Cursor  •  GitHub Copilot  •  Lovable  •  n8n  •  PostHog  •  Figma")

out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Probably_Genetic_SPM.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
