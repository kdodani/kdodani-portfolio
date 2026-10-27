from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)
    section.page_width    = Inches(8.5)
    section.page_height   = Inches(11)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=4, space_after=1, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=3, after=0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=0, after=0.5)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=0, after=0)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY
    return p

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("SENIOR PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Money Movement Flows  •  Funnel Conversion  •  Design Craft  •  Regulated Experiences"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product manager with 5 years owning transaction-heavy, regulated journeys end to end — submission, "
    "validation, eligibility, and payment — where customers decide whether they trust the product. Finds the "
    "friction and failure points in the funnel, fixes them with design and engineering, and measures the lift: "
    "paper 56% → 19%, portal 39% → 55%, invalid submissions −40%. $1.2M+ funded, including $580K to scale a "
    "billing rebuild.",
    size=9.5,
    space_before=3,
    space_after=1,
)
body(
    doc,
    "Leads 0→1 launches and holds a high bar for craft: cut an 11-screen flow to 5, built reusable "
    "error-validation patterns with designers, and shipped with compliance, policy, and engineering partners "
    "so every experience is compliant without being clunky. Instruments before launch and iterates with "
    "experiments. Master of Business Analytics (Python, SQL, machine learning); BBA in Marketing.",
    size=9.5,
    space_before=0,
    space_after=2,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager", "WorkSafeBC", "Apr 2024 – Present")
sublabel(
    doc,
    "British Columbia’s workers’ compensation insurer — billing, eligibility, authorization, and identity.",
)

bullet(doc,
    "Own the roadmap for billing, eligibility, authorization, and identity with engineering, design, operations, "
    "policy, compliance, and executives. Built the business cases that secured $1.2M+ in funding.")
bullet(doc,
    "Rebuilt the invoice-submission journey, where providers send billing to be paid. Analyzed 3.7M line items: "
    "25 of 3,795 organizations drove 50% of paper (~280K lines). Cut 11 screens → 5 with multi-claim entry. "
    "Paper 56% → 19%; portal 39% → 55%; 60% digital adoption in six months; support calls −30%; NPS +25; "
    "~$200K annual savings. Secured $580K to scale.")
bullet(doc,
    "Removed failure points before money moves: in-product validation against fee schedules, rates, and service "
    "codes, and contract eligibility exposed on the API so organizations only see what they can bill. Invalid "
    "submissions −40%; bad claims stopped at submit, not in cleanup.")
bullet(doc,
    "Held a high bar for craft with design. Co-built reusable error-validation patterns, reordered the flow so "
    "worker details are entered once and validated up front, and made usability fixes (claim-number tooltip, "
    "removed low-value subtext, required follow-up fields). Negotiated new components only where reuse justified them.")
bullet(doc,
    "Let research and data set priorities. Discovery showed office admins, not clinicians, were the real users, "
    "so I segmented by workflow and fixed survey bias. Request-aging data (most resolved within 3 days) led me to "
    "deprioritize a \"ping\" feature and test a call-back feature instead.")
bullet(doc,
    "Instrumented PostHog (submit, error, drop-off) before launch. Ran A/B tests on onboarding and "
    "discoverability, and dark → soft → full launches — 322K existing digital lines/year stayed live while a "
    "high-volume segment activated.")
bullet(doc,
    "Partnered with policy and risk on an $80M program where 79% of 139K requests were always approved. Declined "
    "auto-approval — it would have weakened review on the largest spend — and shipped tracking, validation, and "
    "decision rationale so humans kept the decision.")
bullet(doc,
    "Led a 0→1 API with engineering and 50 EMR vendors, TELUS, and government partners after users would not "
    "leave their system of record (API gateway, OAuth 2.0 / mTLS). $80K proof; roadmap for ~70% of the "
    "ecosystem. Migrated 5,600 users across 500 organizations and 62K+ annual referrals to new identity — 0 downtime.")
bullet(doc,
    "Built Referral Radar (LLM retrieval across 20–75 documents per case) with evidence traces and a human "
    "decision-maker; synthetic-data path when production data was blocked. 2nd place, InnoFest 2025.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager (promoted from Business Analyst)", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Built membership, churn, and engagement products that identified 11.7% annual churn; RFM segmentation "
    "retained 78% of high-risk customers.")
bullet(doc,
    "Shipped NLP email classification into a live routing workflow, reducing manual inbox processing.")
bullet(doc,
    "Led discovery for a mobile inspections platform through field interviews, workflow mapping, prototyping, and "
    "usability testing before engineering investment; owned requirements and roadmap in a compliance-constrained operation.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting (UBC)", "Sep 2020 – Oct 2021")
bullet(doc,
    "Designed a conversion-focused acquisition system for Open Primaries that cut bounce rate 37% while doubling subscriptions.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant — Growth & Analytics", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Built acquisition and forecasting systems for 75+ digital properties: organic performance ~150%, SEO "
    "forecast accuracy ~85%.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT, space_before=0, space_after=1)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")
body(
    doc,
    "ProductCamp BC speaker (2025): “How Legacy Industries Can Adopt a Product Mindset.” Mentor, ProductBC.",
    size=9,
    space_before=0,
    space_after=1,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

skills_row(doc, "Product",
    "Strategy  •  Roadmaps  •  0→1 Launches  •  Funnel Conversion  •  Transaction Success  •  Executive Business Cases")
skills_row(doc, "Experience & design",
    "Design Partnership  •  Error-State UX  •  Usability Testing  •  Prototyping  •  User Interviews  •  Journey Mapping")
skills_row(doc, "Analytics & experimentation",
    "PostHog  •  Funnels & Cohorts  •  A/B Testing  •  Staged Launches  •  SQL  •  Python  •  RFM")
skills_row(doc, "Regulated products & platform",
    "Eligibility & Validation  •  Risk vs. Automation  •  Policy & Compliance  •  REST APIs  •  OAuth / mTLS  •  Vendor Integrations")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Robinhood.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
