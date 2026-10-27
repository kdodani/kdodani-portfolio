"""Shared Word helpers for interview playbooks. Same palette as the resume generators."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK = RGBColor(0x1C, 0x19, 0x17)
MID = RGBColor(0x5B, 0x21, 0xB6)
ACCENT = RGBColor(0x6D, 0x28, 0xD9)
BODY = RGBColor(0x44, 0x40, 0x3C)
LIGHT = RGBColor(0x78, 0x71, 0x6C)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
BG_HEX = "FAF9F6"


def new_doc():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
    body_elem = doc.element.body
    doc_elem = body_elem.getparent()
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), BG_HEX)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement("w:displayBackgroundShape")
    settings.insert(0, disp)
    return doc


def _space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)


def rule(doc):
    p = doc.add_paragraph()
    _space(p, 4, 4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "6D28D9")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def title_block(doc, kicker, title, subtitle):
    p = doc.add_paragraph()
    _space(p, 0, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = ACCENT

    p2 = doc.add_paragraph()
    _space(p2, 0, 4)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(title)
    r2.bold = True
    r2.font.size = Pt(18)
    r2.font.color.rgb = DARK

    p3 = doc.add_paragraph()
    _space(p3, 0, 8)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(subtitle)
    r3.font.size = Pt(10)
    r3.font.color.rgb = LIGHT
    rule(doc)


def h1(doc, text):
    p = doc.add_paragraph()
    _space(p, 12, 4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = ACCENT
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    _space(p, 10, 3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = DARK
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    _space(p, 8, 2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = INDIGO
    return p


def body(doc, text, italic=False, size=10, before=2, after=4):
    p = doc.add_paragraph()
    _space(p, before, after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = BODY
    r.italic = italic
    return p


def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    _space(p, 1, 2)
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = BODY
    return p


def labeled(doc, label, text, size=10):
    p = doc.add_paragraph()
    _space(p, 1, 2)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(size)
    rl.font.color.rgb = MID
    rs = p.add_run(text)
    rs.font.size = Pt(size)
    rs.font.color.rgb = BODY
    return p


def say(doc, lines, seconds="60–90s"):
    """Spoken answer block. `lines` is a list of (label, text) tuples."""
    cap = doc.add_paragraph()
    _space(cap, 4, 2)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(f"SAY  ·  {seconds}")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = ACCENT
    for label, text in lines:
        if text:
            labeled(doc, label, text, size=10)


def note(doc, text):
    body(doc, text, italic=True, size=9.5, before=1, after=4)
