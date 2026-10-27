from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

DARK = RGBColor(0x1C, 0x19, 0x17)
MID = RGBColor(0x5B, 0x21, 0xB6)
ACCENT = RGBColor(0x6D, 0x28, 0xD9)
BODY = RGBColor(0x44, 0x40, 0x3C)
LIGHT = RGBColor(0x78, 0x71, 0x6C)
BG_HEX = "FAF9F6"

LINKEDIN_URL = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(document, hex_color):
    body_elem = document.element.body
    doc_elem = body_elem.getparent()
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), hex_color)
    doc_elem.insert(0, bg)
    settings = document.settings.element
    disp = OxmlElement("w:displayBackgroundShape")
    settings.insert(0, disp)


set_page_background(doc, BG_HEX)


def set_run_font(run, name="Calibri"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1.0


def add_horizontal_rule(document, color=ACCENT, thickness=6):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=1, after=1)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(thickness))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def heading(document, text, size=10.5, space_before=6, space_after=1):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = ACCENT
    set_run_font(run)
    return p


def subheading(document, role, org, period, size=10.5):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=4, after=0)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    set_run_font(r1)
    r2 = p.add_run(f"    {period}")
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    set_run_font(r2)
    return p


def sublabel(document, text, size=9, space_before=2, space_after=0):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = True
    run.bold = False
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
    set_run_font(run)
    return p


def bullet(document, text, size=9.5):
    p = document.add_paragraph(style="List Bullet")
    clear_paragraph_space(p, before=0.5, after=0.5)
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    set_run_font(run)
    return p


def body(document, text, size=9.5, color=BODY, space_before=1, space_after=1, italic=False):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    set_run_font(run)
    return p


def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:cs"), "Calibri")
    rPr.append(rFonts)
    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)
    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    set_run_font(run)
    return run


normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(9.5)
normal.font.color.rgb = BODY
rPr = normal.element.get_or_add_rPr()
rFonts = rPr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = OxmlElement("w:rFonts")
    rPr.append(rFonts)
rFonts.set(qn("w:ascii"), "Calibri")
rFonts.set(qn("w:hAnsi"), "Calibri")
rFonts.set(qn("w:cs"), "Calibri")

list_style = doc.styles["List Bullet"]
list_style.font.size = Pt(9.5)
list_style.font.name = "Calibri"
list_style.paragraph_format.space_before = Pt(0)
list_style.paragraph_format.space_after = Pt(0)
list_style.paragraph_format.line_spacing = 1.0

# ── Header ────────────────────────────────────────────────────────────────────
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=0)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK
set_run_font(r)

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=0)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "Relocating to New York City  |  US work authorization  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=1, after=0)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER — AI AGENTS FOR HEALTHCARE BILLING")
rt.bold = True
rt.font.size = Pt(11.5)
rt.font.color.rgb = ACCENT
set_run_font(rt)

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=0, after=1)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Claims & Denial Prevention  ·  Payer–Provider Workflows  ·  Agents with Human Review  ·  0-to-1 and Scale"
)
rp.font.size = Pt(9)
rp.font.color.rgb = MID
set_run_font(rp)

add_horizontal_rule(doc)

# ── Summary ───────────────────────────────────────────────────────────────────
body(
    doc,
    "Product manager with 5 years in product, the last two and a half inside a workers’ compensation payer, on the "
    "side of the claim where denials start: the rules that reject a bad claim, the authorization decisions, and the "
    "faxes, paper, and provider calls that follow when a provider cannot see why. Owns a product surface from customer "
    "discovery to launch to the number it moves (error rate, paper share, cost, turnaround). Pulls own data in SQL "
    "and PostHog, prototypes in Cursor, Copilot, and Lovable, and sets the human-review boundary before an AI feature "
    "ships. Master of Business Analytics.",
    size=9.5,
    space_before=2,
    space_after=2,
)

add_horizontal_rule(doc)

# ── Experience ────────────────────────────────────────────────────────────────
heading(doc, "Experience", space_before=4)

subheading(
    doc,
    "Product Manager — Claims, Billing & Provider Operations",
    "WorkSafeBC",
    "Apr 2024 – Present",
)
sublabel(doc, "Workers’ compensation payer  ·  coverage rules, treatment authorization, and provider billing across 3,795 organizations")

sublabel(doc, "Denial prevention and claims rules", space_before=2)
bullet(
    doc,
    "Encoded fee schedules, contracted rates, service codes, and contract eligibility in the product so invalid "
    "services, claims, and line items fail at submit, not in operations cleanup. The invoice API returns only the "
    "services an organization is contracted to bill. Invalid submissions −40%.",
)
bullet(
    doc,
    "Shipped field-level validation after bad inputs reached production, including a wrong payee number on a live "
    "request. Incomplete or inaccurate identifiers fail at submit with a specific reason, and operations receives "
    "only requests that have already cleared that check.",
)
bullet(
    doc,
    "Digitized continued-treatment authorization on an $80M physiotherapy program after analyzing 139K historical "
    "requests (79% always approved; 27% same-day). Providers were phoning and faxing because they could not see the "
    "decision or the reason. Shipped tracking, validation, and the decision rationale so status no longer needs a call. "
    "Rejected straight-through approval: it would have rubber-stamped spend.",
)

sublabel(doc, "Customer discovery, pilots, and channel strategy", space_before=2)
bullet(
    doc,
    "Analyzed 3.7M invoice line items in SQL. 25 of 3,795 organizations drove half of remaining paper (~280K lines) "
    "because the product capped them at one claim, 11 screens, and ~32 lines. Rebuilt to 5 screens, multi-claim, no line "
    "ceiling. Paper 56% → 19%. Portal 39% → 55%. ~$50K lower processing cost. $580K funded to scale it.",
)
bullet(
    doc,
    "Ran the rollout like a pilot: dark launch to 10 named providers, then soft and hard launch with Healthcare "
    "Operations, with existing volume left intact. Instrumented submit, error, and drop-off in PostHog on 322K digital "
    "lines a year. The first month showed the new path going unused; sat with those providers, found a control they "
    "could not see, and fixed discoverability before widening the rollout.",
)
bullet(
    doc,
    "Chose build or partner by where each provider’s data already lived: rebuilt the portal for clinics billing "
    "200–1,900 lines a month, partnered on a Blueprint bulk feed for hearing-aid practices, and built a direct API into "
    "OSCAR EMR for physicians who will not leave the chart. $80K to prove it. Roadmap covering ~70% of the provider "
    "ecosystem across 50 EMR vendors and TELUS.",
)
bullet(
    doc,
    "The reports that decide a claim were 52,000 faxes a year, truncated by a 2001 billing spec, with physicians "
    "guessing codes. The EMR path carries the full clinical note because a reviewer cannot make an accurate call on a "
    "cut-off record, and neither can an agent.",
)

sublabel(doc, "AI, evals, and the human boundary", space_before=2)
bullet(
    doc,
    "Built Referral Radar in a two-week sprint: LLM retrieval over 20–75 claim documents per case to surface the program "
    "a patient qualified for, with an evidence trace back to the source. Set the bar as a trace a case manager can check, "
    "not an automation rate. Production records were blocked, so the test set was synthetic and checked with case managers. "
    "The person keeps the decision. 2nd of 36, InnoFest 2025.",
)
bullet(
    doc,
    "Prototypes in Cursor, GitHub Copilot, and Lovable to test an idea with providers and case managers before asking "
    "engineering to build it. Migrated 5,600 healthcare professionals across 500 organizations on a portal handling "
    "62,000+ referrals a year: zero downtime, every active account moved on the deadline.",
)

subheading(doc, "Product Manager (promoted from Business Analyst)", "Technical Safety BC", "Oct 2021 – Apr 2024")
sublabel(doc, "Revenue recovery: collections outreach and reactivation", space_before=1)
bullet(
    doc,
    "Controlled test of RFM-targeted calls against the existing call strategy over the same two weeks, success defined "
    "before the calls went out: $45,742 collected versus $25,597 (78% more) on 62 calls versus 48. Recommendation was to "
    "scale the RFM treatment.",
)
bullet(
    doc,
    "Tested a three-step email sequence on expired permits: 135 clients, about CAD 46.7K recovered, 23% paid from the "
    "first email. 2022 permit volume finished 8% above budget.",
)
bullet(
    doc,
    "Built the churn and engagement dashboards the team ran from (11.7% annual churn across ~102K permits). Shipped NLP "
    "routing that took operations email triage off a human inbox. Led field discovery for a mobile inspections product "
    "through prototypes and usability tests.",
)

subheading(doc, "Strategy Consultant", "180 Degrees Consulting — UBC", "Sep 2020 – Oct 2021")
bullet(
    doc,
    "Customer research and market strategy for nonprofit and clean-energy clients. Acquisition work for Open Primaries "
    "cut bounce rate 37% and doubled subscriptions.",
)

subheading(doc, "SEO Consultant", "Convertus", "May 2018 – Jul 2019")
bullet(
    doc,
    "Analytics-led growth for 75+ client sites (~150% organic performance) and a forecast that predicted results at ~85% accuracy.",
)

add_horizontal_rule(doc)

# ── Education ─────────────────────────────────────────────────────────────────
heading(doc, "Education", space_before=4)
subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(
    doc,
    "Python  ·  SQL  ·  Machine learning  ·  Optimization  ·  Product analytics",
    size=9,
    color=LIGHT,
    space_before=0,
    space_after=0,
)
subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ── Leadership ────────────────────────────────────────────────────────────────
heading(doc, "Leadership", space_before=4)
bullet(
    doc,
    "Speaker, ProductCamp BC (2025) — “How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery.” "
    "Mentor, ProductBC — discovery, strategy, and career development for aspiring product managers.",
)

add_horizontal_rule(doc)

# ── Skills ────────────────────────────────────────────────────────────────────
heading(doc, "Skills", space_before=4)


def skills_row(document, label, skills):
    p = document.add_paragraph()
    clear_paragraph_space(p, before=1, after=0)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = MID
    set_run_font(rl)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY
    set_run_font(rs)


skills_row(
    doc,
    "Product",
    "Product strategy and roadmaps  ·  Customer discovery  ·  Pilots and go-lives  ·  Specs and acceptance criteria  ·  0-to-1 and scale",
)
skills_row(
    doc,
    "Healthcare billing",
    "Claims rules  ·  Fee schedules and codes  ·  Eligibility  ·  Authorization  ·  Denial prevention  ·  Portal, fax, API, and EMR intake  ·  Collections",
)
skills_row(
    doc,
    "Data and AI",
    "SQL  ·  Python  ·  PostHog  ·  Dashboards  ·  LLM retrieval with evidence traces  ·  Eval design  ·  Human-in-the-loop  ·  Cursor, Copilot, Lovable",
)

out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Amperos.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
