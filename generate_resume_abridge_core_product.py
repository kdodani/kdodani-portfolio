from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)

# ── Portfolio color system (matches globals.css + design-system.md) ───────────
# Background: #faf9f6  (warm off-white — body bg)
# Headings:   #1c1917  (stone-900 — near-black warm)
# Body text:  #44403c  (stone-700)
# Muted:      #78716c  (stone-500)
# Accent:     #6D28D9  (--accent-hero-1 purple)
# Mid-accent: #5B21B6  (--accent-strong)
# Sub-accent: #4F46E5  (--accent-hero-2 indigo)
DARK    = RGBColor(0x1c, 0x19, 0x17)   # stone-900 — headings
MID     = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple — section labels
ACCENT  = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple — rules / name accent
BODY    = RGBColor(0x44, 0x40, 0x3c)   # stone-700 — body text
LIGHT   = RGBColor(0x78, 0x71, 0x6c)   # stone-500 — muted / metadata
BG_HEX  = "FAF9F6"                     # warm off-white page background

LINKEDIN_URL  = "https://www.linkedin.com/in/khushdodani"
GITHUB_URL    = "https://github.com/kdodani"
PORTFOLIO_URL = "https://kdodani-portfolio.vercel.app"


def set_page_background(doc, hex_color):
    """Set the document page background color."""
    body_elem = doc.element.body
    doc_elem  = body_elem.getparent()
    bg = OxmlElement('w:background')
    bg.set(qn('w:color'), hex_color)
    doc_elem.insert(0, bg)
    settings = doc.settings.element
    disp = OxmlElement('w:displayBackgroundShape')
    settings.insert(0, disp)

set_page_background(doc, BG_HEX)

def clear_paragraph_space(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)

def add_horizontal_rule(doc, color=RGBColor(0x6D, 0x28, 0xD9), thickness=6):
    """Insert a thin colored paragraph border as a visual divider."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(thickness))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), f'{color[0]:02X}{color[1]:02X}{color[2]:02X}')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading(doc, text, size=10.5, bold=True, color=RGBColor(0x6D, 0x28, 0xD9), space_before=8, space_after=2, caps=True):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def subheading(doc, role, org, period, size=10):
    """Role | Org   ·   Period  — two runs on one line."""
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=5, after=1)
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(f"{role}  |  {org}")
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = DARK
    r2 = p.add_run(f"    {period}")
    r2.bold = False
    r2.font.size = Pt(9)
    r2.font.color.rgb = LIGHT
    return p

def sublabel(doc, text, size=9, italic=True, color=RGBColor(0x4F, 0x46, 0xE5), space_before=0, space_after=1):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.italic = italic
    run.bold = not italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text, size=9.5, indent=Inches(0.2), hanging=Inches(0.2)):
    p = doc.add_paragraph(style='List Bullet')
    clear_paragraph_space(p, before=1, after=1)
    p.paragraph_format.left_indent   = indent + hanging
    p.paragraph_format.first_line_indent = -hanging
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = BODY
    return p

def body(doc, text, size=9.5, color=BODY, space_before=2, space_after=2, italic=False):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=space_before, after=space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.italic = italic
    return p

def add_hyperlink(paragraph, text, url, font_size=Pt(9), color=LIGHT):
    """Append a clickable hyperlink run that matches surrounding muted contact styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color_elem = OxmlElement("w:color")
    color_elem.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    rPr.append(color_elem)

    sz_val = str(int(font_size.pt * 2))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), sz_val)
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), sz_val)
    rPr.append(szCs)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def contact_run(paragraph, text, font_size=Pt(9)):
    run = paragraph.add_run(text)
    run.font.size = font_size
    run.font.color.rgb = LIGHT
    return run

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
p_name = doc.add_paragraph()
clear_paragraph_space(p_name, before=0, after=2)
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_name.add_run("KHUSHBOO DODANI")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = DARK

p_contact = doc.add_paragraph()
clear_paragraph_space(p_contact, before=0, after=2)
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_run(p_contact, "San Francisco Bay Area  |  khushdodani@gmail.com  |  ")
add_hyperlink(p_contact, "LinkedIn", LINKEDIN_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "GitHub", GITHUB_URL)
contact_run(p_contact, "  |  ")
add_hyperlink(p_contact, "Portfolio", PORTFOLIO_URL)

p_title = doc.add_paragraph()
clear_paragraph_space(p_title, before=2, after=2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = p_title.add_run("PRODUCT MANAGER")
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)   # accent-hero-1 purple

p_pills = doc.add_paragraph()
clear_paragraph_space(p_pills, before=1, after=4)
p_pills.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pills.add_run(
    "Clinical Documentation Workflows  •  EHR-Embedded Experiences  •  "
    "Ambient & AI-Generated Documentation  •  Clinician Adoption"
)
rp.font.size = Pt(9)
rp.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)   # accent-strong purple

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
body(
    doc,
    "Product Manager building clinician-facing documentation products and AI-assisted workflows in regulated "
    "healthcare. Focused on reducing documentation burden and improving note quality within the tools clinicians "
    "already use. Ships AI-generated clinical documentation and artifacts with evidence traceability and "
    "explicit human-in-the-loop boundaries. Master of Business Analytics; BBA in Marketing.",
    size=9.5,
    space_before=4,
    space_after=4,
)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EXPERIENCE
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Experience")

# ── WorkSafeBC ────────────────────────────────────────────────────────────────
subheading(doc, "Product Manager — Clinician Experience & Clinical Documentation", "WorkSafeBC", "Apr 2024 – Present")

sublabel(doc, "Clinical Documentation & Note Quality")
bullet(doc,
    "Own the clinician-facing products physicians and care teams use to document care, request treatment, and "
    "close out administrative work after the patient leaves.")
bullet(doc,
    "Treated documentation completeness as a product quality problem, not a training problem: Form 8/11 clinical "
    "notes were being truncated by a 2001 billing specification that had frozen the form since 2015, pushing "
    "52,000 faxes per year and roughly 28M keystrokes of manual re-keying into the system.")
bullet(doc,
    "Shipped front-end validation and a simplified documentation flow that improved the quality of what arrived "
    "downstream, decreased support calls by 30%, and increased clinician satisfaction (NPS +25) — the difference "
    "between a note that is technically submitted and one a reviewer can actually act on.")

sublabel(doc, "EHR Integration & In-Workflow Experience")
bullet(doc,
    "Ran 15 physician interviews, usability sessions, and contextual observation of case managers and clinic "
    "office administrators, producing the core conviction for the roadmap: GPs will not leave their EMR for a "
    "separate application. Redirected the program from portal improvements to EMR-native submission on that "
    "finding.")
bullet(doc,
    "Led 0-to-1 discovery and delivery for point-of-care clinical documentation submitted directly from OSCAR "
    "(community GP EMR) instead of a side portal — POC validated March 2025, proving clinical reporting could be "
    "decoupled from the legacy billing channel it had been trapped in.")
bullet(doc,
    "Made the embedded-vs-standalone tradeoffs explicit with architects: JSON for POC speed against FHIR as the "
    "documented upgrade path, Azure APIM as a reusable facade against point-to-point integrations, OAuth 2.0 / "
    "mTLS for the clinical channel.")
bullet(doc,
    "Secured $80K in Innovation Council funding and a vendor-agnostic integration roadmap covering roughly 70% "
    "of the provider ecosystem, coordinating 50 EMR vendors plus TELUS, ForgeRock, and government partners.")

sublabel(doc, "AI-Generated Clinical Documentation & Clinician Trust")
bullet(doc,
    "Built Referral Radar, an LLM/RAG prototype that reads 20–75 unstructured documents per case — intake notes, "
    "medical reports, counselling plans — to generate summaries, surface supporting evidence, and recommend "
    "referrals, reducing manual review time and enabling faster, evidence-backed decisions; 2nd place of 36 "
    "entries at InnoFest 2025.")
bullet(doc,
    "Built toward ambient documentation from both directions — structured capture at the point of care in the "
    "EMR, and LLM summarization of unstructured clinical records with linked evidence.")
bullet(doc,
    "Rejected a ranked recommendation list in favour of evidence traces back to source documents, because "
    "reviewers will not act on guidance they cannot verify. Set the evaluation bar on explainability and "
    "traceability rather than automation rate.")
bullet(doc,
    "Refused auto-approve on treatment extensions after analysis of 139K historical requests showed 79% were "
    "always approved — on an $80M spend category, that aggregate was the argument against automating, not for "
    "it. Kept clinical judgment with the reviewer; AI retrieves, humans decide.")
bullet(doc,
    "Unblocked a stalled two-week AI sprint when production claim data was unavailable by generating synthetic "
    "case data and validating it with case managers before build.")
bullet(doc,
    "Prototype personally with AI tooling — Lovable and GitHub Copilot for concept-to-validation, n8n for "
    "multi-agent workflows — compressing discovery-to-engineering timelines before build investment.")

sublabel(doc, "Adoption, Retention & Measurement")
bullet(doc,
    "Rebuilt high-volume clinician submission after segmentation showed 25 of 3,795 organizations generated 50% "
    "of remaining paper volume (~280K blocked line items): cut the journey from 11 screens to 5, shifted paper "
    "from 56% to 19% of line items, and grew digital channel share from 39% to 55% — securing $580K in funding.")
bullet(doc,
    "Instrumented PostHog funnels before launch and ran named-cohort dark, soft, and hard launches so adoption "
    "and clinician drop-off were visible within days rather than after a big-bang cutover, protecting existing "
    "digital volume while a new high-volume segment activated.")
bullet(doc,
    "Led a CIAM migration of 5,600 healthcare professionals across 500 organizations on a portal carrying 62,000+ "
    "annual referrals — 0 service disruptions and 100% of active accounts transitioned on deadline.")

# ── Technical Safety BC ───────────────────────────────────────────────────────
subheading(doc, "Product Manager", "Technical Safety BC", "Oct 2021 – Apr 2024")
bullet(doc,
    "Shipped NLP-powered email classification into a live operational workflow to automate routing and triage of "
    "incoming requests — production machine learning in a regulated environment, and the predecessor to later "
    "LLM work.")
bullet(doc,
    "Led discovery for a mobile remote inspections platform through contextual research with field inspectors "
    "and usability testing before development investment.")
bullet(doc,
    "Built operational analytics identifying an 11.7% average annual churn rate and an RFM segmentation model "
    "that drove targeted engagement campaigns retaining 78% of high-risk customers.")

# ── 180 Degrees ───────────────────────────────────────────────────────────────
subheading(doc, "Strategy Consultant", "180 Degrees Consulting — University of British Columbia", "Sep 2020 – Oct 2021")
bullet(doc,
    "Built product and market strategies for nonprofit and clean energy clients; designed a data-driven "
    "acquisition strategy for Open Primaries that cut bounce rate 37% while doubling subscriptions.")

# ── Convertus ─────────────────────────────────────────────────────────────────
subheading(doc, "SEO Consultant", "Convertus", "May 2018 – Jul 2019")
bullet(doc,
    "Built analytics-driven growth strategies for 75+ client websites, increasing organic performance ~150%, "
    "and a forecasting tool predicting SEO performance with ~85% accuracy.")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# EDUCATION
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Education")

subheading(doc, "Master of Business Analytics", "University of British Columbia", "")
body(doc, "Python  •  SQL  •  Machine Learning  •  Optimization  •  Product Analytics", size=9, color=LIGHT)

subheading(doc, "Bachelor of Business Administration (Marketing)", "British Columbia Institute of Technology", "")

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# LEADERSHIP
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Leadership")

sublabel(doc, "Speaker — ProductCamp BC (2025)", italic=False, size=9.5, color=DARK)
body(doc,
    '"How Legacy Industries Can Adopt a Product Mindset to Improve Service Delivery."',
    size=9.5, italic=True, space_before=0)

sublabel(doc, "Mentor — ProductBC", italic=False, size=9.5, color=DARK, space_before=4)
body(doc,
    "Mentor aspiring Product Managers in product discovery, strategy, and career development.",
    size=9.5, space_before=0)

add_horizontal_rule(doc)

# ══════════════════════════════════════════════════════════════════════════════
# SKILLS
# ══════════════════════════════════════════════════════════════════════════════
heading(doc, "Skills")

def skills_row(doc, label, skills):
    p = doc.add_paragraph()
    clear_paragraph_space(p, before=2, after=2)
    rl = p.add_run(f"{label}:  ")
    rl.bold = True
    rl.font.size = Pt(9)
    rl.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    rs = p.add_run(skills)
    rs.font.size = Pt(9)
    rs.font.color.rgb = BODY

skills_row(doc, "Product",
    "Product Strategy  •  Roadmap Ownership  •  0-to-1 Products  •  Prioritization  •  "
    "Adoption & Retention Metrics  •  Change Management  •  Executive Influence")
skills_row(doc, "Clinical Workflows",
    "Clinical Documentation  •  Note Quality & Completeness  •  Point-of-Care Workflows  •  "
    "EMR/EHR Integration  •  Provider Experience  •  Regulated Healthcare")
skills_row(doc, "AI / ML",
    "LLM Applications  •  RAG  •  Multi-Agent Workflows  •  Document Summarization  •  "
    "Evidence Traceability  •  Human-in-the-Loop Design  •  Output Evaluation  •  NLP")
skills_row(doc, "Technical",
    "REST APIs  •  EMR/EHR Integrations  •  Azure APIM  •  OAuth 2.0 / mTLS  •  JSON vs FHIR  •  SQL  •  "
    "Python  •  PostHog  •  Cursor  •  Lovable  •  n8n")
skills_row(doc, "Research",
    "Physician Interviews  •  Contextual Observation  •  Usability Testing  •  Prototyping  •  Figma")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = r"c:\Users\khush\Desktop\kdodani-portfolio\Khushboo_Dodani_Resume_Abridge_Core_Product.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
